"""Phase 02 — ingestion and X-9xx validation.

The failure mode of a refreshed workbook must be a named issue, never a wrong number.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest

from conftest import event, make_workbook, position
from hc_valuation.engine.inputs import EventType, Status
from hc_valuation.ingest.reader import IngestError, read_workbook
from hc_valuation.ingest.validate import validate


def _issues(path: Path, cfg):
    snap, feed = read_workbook(path, cfg)
    return snap, feed, validate(snap, feed, cfg)


def _by_rule(issues, rule_id):
    return [i for i in issues if i.rule_id == rule_id]


# ---------------------------------------------------------------- real workbook

def test_real_workbook_loads_clean(workbook_path, cfg):
    snap, feed, issues = _issues(workbook_path, cfg)
    assert len(snap.positions) == 100
    assert sum(1 for p in snap.positions if p.status == Status.ACTIVE) == 96
    assert len(feed.events) == 18
    assert feed.quarter_label == cfg.quarter.label
    assert feed.sheet_name == "Q3 2026 Activity"
    assert snap.unknown_columns == () and feed.unknown_columns == ()
    assert issues == []


def test_real_workbook_prior_marks_reconcile(workbook_path, cfg):
    snap, _ = read_workbook(workbook_path, cfg)
    tol = cfg.tolerances.prior_mark_reconciliation_musd
    for p in snap.positions:
        if p.status == Status.ACTIVE:
            assert abs(p.ownership * p.latest_post_money - p.prior_mark) <= tol, p.company
        else:
            assert p.prior_mark == 0, p.company


def test_renamed_activity_sheet_loads_by_pattern(workbook_path, cfg, tmp_path):
    copy = tmp_path / "renamed.xlsx"
    shutil.copy(workbook_path, copy)
    wb = openpyxl.load_workbook(copy)
    wb["Q3 2026 Activity"].title = "Q4 2026 Activity"
    wb.save(copy)

    snap, feed = read_workbook(copy, cfg)
    assert feed.sheet_name == "Q4 2026 Activity"
    assert feed.quarter_label == "Q4 2026"
    assert len(feed.events) == 18 and len(snap.positions) == 100


def test_deleted_post_money_is_a_named_x902(workbook_path, cfg, tmp_path):
    copy = tmp_path / "no_post.xlsx"
    shutil.copy(workbook_path, copy)
    wb = openpyxl.load_workbook(copy)
    ws = wb["Q3 2026 Activity"]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("Post-Money / Deal Value ($M)") + 1
    target = None
    for row in ws.iter_rows(min_row=2):
        if row[hdr.index("Company")].value == "Fernwave":
            target = row[0].row
            ws.cell(row=target, column=col).value = None
    assert target is not None
    wb.save(copy)

    _, _, issues = _issues(copy, cfg)
    x902 = _by_rule(issues, "X-902")
    assert len(x902) == 1
    assert x902[0].company == "Fernwave" and x902[0].row_index == target and x902[0].blocking
    assert "post-money" in x902[0].message


def test_missing_activity_sheet_raises_naming_the_pattern(tmp_path, cfg):
    """A sheet that matches neither the policy regex nor the relaxed 'activity'/'events' rule."""
    path = make_workbook(tmp_path, [position()], [], activity_sheet="Deal Log")
    with pytest.raises(IngestError) as ex:
        read_workbook(path, cfg)
    msg = str(ex.value)
    assert cfg.schema_.activity_sheet_pattern in msg
    assert "Deal Log" in msg


def test_two_activity_sheets_is_an_error(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [])
    wb = openpyxl.load_workbook(path)
    wb.create_sheet("Q4 2026 Activity")
    wb.save(path)
    with pytest.raises(IngestError, match="more than one"):
        read_workbook(path, cfg)


def test_missing_required_column_raises(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [])
    wb = openpyxl.load_workbook(path)
    ws = wb["Portfolio"]
    hdr = [c.value for c in ws[1]]
    ws.delete_cols(hdr.index("Prior Mark ($M)") + 1)
    wb.save(path)
    with pytest.raises(IngestError, match=r"Prior Mark \(\$M\)"):
        read_workbook(path, cfg)


def test_unknown_column_is_recorded_and_warns(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [event(value=150.0, ownership_after=0.1)],
                         extra_portfolio_columns={"Board Seat": "Yes"},
                         extra_activity_columns={"Lead Investor": "Acme Ventures"})
    snap, feed, issues = _issues(path, cfg)
    assert snap.unknown_columns == ("Board Seat",)
    assert feed.unknown_columns == ("Lead Investor",)
    assert snap.positions[0].extra == {"Board Seat": "Yes"}
    assert feed.events[0].extra == {"Lead Investor": "Acme Ventures"}
    x910 = _by_rule(issues, "X-910")
    assert len(x910) == 2
    assert all(not i.blocking for i in x910)
    assert "Board Seat" in x910[0].message and "Lead Investor" in x910[1].message
    assert not any(i.blocking for i in issues)


def test_unknown_column_fails_when_policy_says_fail(tmp_path, cfg):
    from conftest import with_policy
    strict = with_policy(cfg, **{"schema.unknown_column": "fail"})
    path = make_workbook(tmp_path, [position()], [], extra_portfolio_columns={"Board Seat": "Yes"})
    with pytest.raises(IngestError, match="Board Seat"):
        read_workbook(path, strict)


# ---------------------------------------------------------------- X-9xx on the portfolio tab

def test_prior_mark_not_tying_is_x904(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(prior_mark=12.0)], [])   # 0.10 × 100 = 10, not 12
    _, _, issues = _issues(path, cfg)
    x904 = _by_rule(issues, "X-904")
    assert len(x904) == 1 and x904[0].company == "Alpha" and x904[0].blocking
    assert "10.00" in x904[0].message


def test_prior_mark_within_tolerance_is_fine(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(prior_mark=10.04)], [])
    _, _, issues = _issues(path, cfg)
    assert not _by_rule(issues, "X-904")


def test_non_active_company_with_a_mark_is_x904(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(status="Acquired", prior_mark=3.0, ownership=0.0, realized=20.0)], [])
    _, _, issues = _issues(path, cfg)
    assert [i.rule_id for i in issues] == ["X-904"]


def test_ownership_above_one_hundred_is_x903(tmp_path, cfg):
    """150 is neither a fraction nor percentage points (SPEC §2.3: > 100 blocks). The prior
    mark ties (150 × 100), so the only issue is the refusal, reported once."""
    path = make_workbook(tmp_path, [position(ownership=150)], [])
    _, _, issues = _issues(path, cfg)
    assert [i.rule_id for i in issues] == ["X-903"]
    assert issues[0].company == "Alpha" and issues[0].blocking and "150" in issues[0].message


def test_ownership_in_percentage_points_is_read_and_flagged_x916(tmp_path, cfg):
    """1.5 on an ownership column is 1.5 points (SPEC §2.3): read as 0.015 with X-916, and the
    prior mark is then checked against the corrected fraction."""
    path = make_workbook(tmp_path, [position(ownership=1.5, prior_mark=1.5)], [])
    snap, _, issues = _issues(path, cfg)
    assert snap.positions[0].ownership == pytest.approx(0.015)
    x916 = _by_rule(issues, "X-916")
    assert len(x916) == 1 and not x916[0].blocking and "percentage points" in x916[0].message
    assert not any(i.blocking for i in issues)


def test_duplicate_company_row_is_x906(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(), position()], [])
    _, _, issues = _issues(path, cfg)
    assert [i.rule_id for i in issues] == ["X-906", "X-906"] and all(i.blocking for i in issues)   # every copy blocks


def test_unknown_status_raises(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(status="Zombie")], [])
    with pytest.raises(IngestError, match="Zombie"):
        read_workbook(path, cfg)


# ---------------------------------------------------------------- X-9xx on the activity tab

def test_event_on_acquired_company_is_x907(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(status="Acquired", prior_mark=0.0, ownership=0.0, realized=20.0)],
                         [event(value=150.0, ownership_after=0.1)])
    _, _, issues = _issues(path, cfg)
    x907 = _by_rule(issues, "X-907")
    assert len(x907) == 1 and x907[0].company == "Alpha" and x907[0].blocking
    assert "Acquired" in x907[0].message


def test_unknown_event_type_is_x909(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [event("SPAC Merger", value=400.0)])
    _, _, issues = _issues(path, cfg)
    x909 = _by_rule(issues, "X-909")
    assert len(x909) == 1 and "SPAC Merger" in x909[0].message and x909[0].blocking


def test_event_for_company_not_in_book_is_x901(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [event(company="Ghost", value=1.0, ownership_after=0.1)])
    _, _, issues = _issues(path, cfg)
    assert [i.rule_id for i in issues] == ["X-901"]


def test_event_outside_window_is_x905_non_blocking(tmp_path, cfg):
    from datetime import date
    path = make_workbook(tmp_path, [position()], [event(date=date(2026, 6, 15), value=150.0, ownership_after=0.1)])
    _, _, issues = _issues(path, cfg)
    assert [i.rule_id for i in issues] == ["X-905"] and not issues[0].blocking


@pytest.mark.parametrize("et,kw,fragment", [
    (EventType.PRICED_ROUND, dict(value=150.0), "ownership after"),
    (EventType.IPO, dict(ownership_after=0.05), "market cap"),
    (EventType.ACQ_CLOSED, dict(proceeds=10.0), "deal value"),
    (EventType.ACQ_ANNOUNCED, dict(), "deal value"),
    (EventType.SECONDARY, dict(proceeds=1.0), "ownership after"),
])
def test_missing_mandatory_field_is_x902(tmp_path, cfg, et, kw, fragment):
    path = make_workbook(tmp_path, [position()], [event(et, **kw)])
    _, _, issues = _issues(path, cfg)
    x902 = _by_rule(issues, "X-902")
    assert x902 and all(i.company == "Alpha" for i in x902)
    assert any(fragment in i.message for i in x902), [i.message for i in x902]


def test_event_ownership_after_out_of_range_is_x903(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [event(value=150.0, ownership_after=120)])
    _, _, issues = _issues(path, cfg)
    assert [i.rule_id for i in issues] == ["X-903"]
    assert issues[0].blocking and "120" in issues[0].message


def test_string_cells_are_coerced(tmp_path, cfg):
    """Hand-edited cells arrive as text: '$120.0', '12%', '2026-08-15' must still parse."""
    path = make_workbook(tmp_path, [position(latest_post_money="$100.0", cash="12", prior_mark=10.0, sheet_moic=2.0, sheet_runway=24.0)],
                         [event(date="2026-08-15", value="$150.0", ownership_after="0.10")])
    snap, feed = read_workbook(path, cfg)
    assert snap.positions[0].latest_post_money == 100.0 and snap.positions[0].cash == 12.0
    e = feed.events[0]
    assert e.value == 150.0 and e.ownership_after == 0.10 and e.date.isoformat() == "2026-08-15"


# ---------------------------------------------------------------- normalization (SPEC §2): what was read as what

def _edit(path: Path, fn) -> Path:
    """Open a synthetic workbook, let `fn(wb)` dirty it, save it back."""
    wb = openpyxl.load_workbook(path)
    fn(wb)
    wb.save(path)
    return path


def _act(wb):
    return wb["Q3 2026 Activity"]


def test_event_type_typo_resolves_to_canonical_with_x912(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [event("Priced Equtiy Round", value=150.0, ownership_after=0.1)])
    snap, feed, issues = _issues(path, cfg)
    e = feed.events[0]
    assert e.event_type == EventType.PRICED_ROUND.value and e.extra["raw_event_type"] == "Priced Equtiy Round"
    x912 = _by_rule(issues, "X-912")
    assert len(x912) == 1 and not x912[0].blocking and "distance=1" in x912[0].message
    assert "'Priced Equtiy Round' → 'Priced Equity Round'" in x912[0].message
    assert not _by_rule(issues, "X-909") and not any(i.blocking for i in issues)


def test_event_type_synonym_resolves_with_x912(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [event("Series B", value=150.0, ownership_after=0.1)])
    _, feed, issues = _issues(path, cfg)
    assert feed.events[0].event_type == EventType.PRICED_ROUND.value
    assert [i.rule_id for i in issues] == ["X-912"] and "synonym" in issues[0].message


@pytest.mark.parametrize("raw", ["Acquisition", "Sale", "loi", "LOI"])
def test_ambiguous_event_type_is_a_blocking_x914_naming_candidates(tmp_path, cfg, raw):
    path = make_workbook(tmp_path, [position()], [event(raw, value=150.0)])
    _, feed, issues = _issues(path, cfg)
    assert feed.events[0].event_type == raw                      # left raw for M-999
    x914 = _by_rule(issues, "X-914")
    assert len(x914) == 1 and x914[0].blocking and x914[0].company == "Alpha"
    assert not _by_rule(issues, "X-909")                         # X-914 already names the problem
    assert x914[0].message.count("Acquisition") + x914[0].message.count("Secondary") + x914[0].message.count("Term Sheet") >= 2


def test_company_name_variants_resolve_with_x913(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(company="Nimbrel Systems")], [
        event(company="NIMBREL SYSTEMS INC.", value=150.0, ownership_after=0.1),
        event(company="Nimbrel Sytems", value=150.0, ownership_after=0.1, date=__import__("datetime").date(2026, 8, 16)),
        event(company="Skylark (formerly Nimbrel Systems)", value=150.0, ownership_after=0.1, date=__import__("datetime").date(2026, 8, 17)),
    ])
    _, feed, issues = _issues(path, cfg)
    assert all(e.company == "Nimbrel Systems" for e in feed.events)
    assert feed.events[0].extra["raw_company"] == "NIMBREL SYSTEMS INC."
    x913 = _by_rule(issues, "X-913")
    assert [m for m in ("suffix", "distance=1", "formerly") if any(m in i.message for i in x913)] == ["suffix", "distance=1", "formerly"]
    assert not _by_rule(issues, "X-901") and not any(i.blocking for i in issues)


def test_near_duplicate_book_names_make_a_typo_ambiguous_x914(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(company="Aravine Labs"), position(company="Aravine Lab")],
                         [event(company="Aravine Labz", value=150.0, ownership_after=0.1)])
    _, _, issues = _issues(path, cfg)
    x914 = _by_rule(issues, "X-914")
    assert len(x914) == 1 and x914[0].blocking and "Aravine Labs" in x914[0].message and "Aravine Lab" in x914[0].message


def test_value_formats_are_coerced_and_recorded_x915(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(latest_post_money="$100.0M", cash="12", prior_mark=10.0, sheet_moic=2.0, sheet_runway=24.0)],
                         [event(date="15-Aug-2026", value="$150,000,000", ownership_after="10%", proceeds="(1.2)")])
    snap, feed, issues = _issues(path, cfg)
    p, e = snap.positions[0], feed.events[0]
    assert p.latest_post_money == 100.0 and p.cash == 12.0
    assert e.date.isoformat() == "2026-08-15" and e.value == pytest.approx(150.0)
    assert e.ownership_after == pytest.approx(0.10) and e.proceeds == -1.2
    x915 = _by_rule(issues, "X-915")
    assert len(x915) == 6 and all(not i.blocking for i in x915)
    x916 = _by_rule(issues, "X-916")
    assert len(x916) == 1 and "dollars" in x916[0].message and not x916[0].blocking
    # "(1.2)" is read as -1.2 (X-915, coercion) — and negative proceeds are then a domain
    # error the row must block on (X-903). Coercion and validation are separate findings.
    blocking = [i for i in issues if i.blocking]
    assert [i.rule_id for i in blocking] == ["X-903"] and "proceeds" in blocking[0].message


def test_units_post_money_in_dollars_and_ownership_in_points_x916(tmp_path, cfg):
    path = make_workbook(tmp_path, [position(ownership=10, prior_mark=10.0)],
                         [event(value=150_000_000, ownership_after=8.5, proceeds=1_200_000)])
    snap, feed, issues = _issues(path, cfg)
    assert snap.positions[0].ownership == pytest.approx(0.10)
    e = feed.events[0]
    assert e.value == pytest.approx(150.0) and e.ownership_after == pytest.approx(0.085) and e.proceeds == pytest.approx(1.2)
    x916 = _by_rule(issues, "X-916")
    assert len(x916) == 4 and all(i.severity.value == "REVIEW" and not i.blocking for i in x916)
    assert not any(i.blocking for i in issues)


def test_arr_growth_above_one_is_a_fraction_not_points(tmp_path, cfg):
    """Growth of 207% is real (the Q3 book carries 3.0); the points heuristic must not touch it."""
    path = make_workbook(tmp_path, [position(arr_growth=2.07)], [])
    snap, _, issues = _issues(path, cfg)
    assert snap.positions[0].arr_growth == 2.07 and issues == []


def test_text_in_a_number_cell_is_a_blocking_x902_naming_the_cell(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [event(value="about 150", ownership_after=0.1)])
    _, feed, issues = _issues(path, cfg)
    assert feed.events[0].value is None
    x902 = [i for i in _by_rule(issues, "X-902") if "not a number" in i.message]
    assert len(x902) == 1 and x902[0].blocking and "about 150" in x902[0].message and "Post-Money" in x902[0].message


def test_unparseable_date_is_an_ingest_error_naming_the_cell(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [event(date="sometime in August", value=150.0, ownership_after=0.1)])
    with pytest.raises(IngestError, match=r"row 2.*'Date'.*sometime in August"):
        read_workbook(path, cfg)


def test_excel_serial_and_day_first_ambiguity_are_x915(tmp_path, cfg):
    from datetime import date
    path = make_workbook(tmp_path, [position()], [
        event(date=46249, value=150.0, ownership_after=0.1),            # 2026-08-15 as a serial
        event(date="08/03/2026", value=150.0, ownership_after=0.1),     # month-first, ambiguous
    ])
    _, feed, issues = _issues(path, cfg)
    assert feed.events[0].date == date(2026, 8, 15) and feed.events[1].date == date(2026, 8, 3)
    x915 = _by_rule(issues, "X-915")
    assert len(x915) == 2 and "excel-serial" in x915[0].message and "ambiguous" in x915[1].message


# ---------------------------------------------------------------- headers (§2.2)

def test_header_aliases_typos_and_reordering_load_with_x911(tmp_path, cfg):
    def dirty(wb):
        ws = wb["Portfolio"]
        hdr = [c.value for c in ws[1]]
        ws.cell(row=1, column=hdr.index("Company") + 1).value = "Portfolio Company"
        ws.cell(row=1, column=hdr.index("Prior Mark ($M)") + 1).value = "Carrying Value"
        ws.cell(row=1, column=hdr.index("Latest Post-Money ($M)") + 1).value = "Latest Post Money ($M)"
        ws.cell(row=1, column=hdr.index("Ownership (FD %)") + 1).value = "ownership (fd %)"
        wa = _act(wb)
        hdr = [c.value for c in wa[1]]
        wa.cell(row=1, column=hdr.index("Post-Money / Deal Value ($M)") + 1).value = "Post Money / Deal Value ($M)"
        wa.cell(row=1, column=hdr.index("Event") + 1).value = "Event Type"
        wa.move_range("A1:B3", cols=len(hdr))      # reorder: Date/Company become the last columns
        wa.delete_cols(1, 2)
    path = _edit(make_workbook(tmp_path, [position()], [event(value=150.0, ownership_after=0.1)]), dirty)
    snap, feed, issues = _issues(path, cfg)
    assert snap.positions[0].company == "Alpha" and snap.positions[0].prior_mark == 10.0
    assert feed.events[0].company == "Alpha" and feed.events[0].value == 150.0
    x911 = _by_rule(issues, "X-911")
    assert len(x911) == 6 and all(not i.blocking for i in x911)
    methods = {m for m in ("alias", "distance=1", "fold") if any(m in i.message for i in x911)}
    assert methods == {"alias", "distance=1", "fold"}
    assert snap.unknown_columns == () and feed.unknown_columns == ()
    assert not any(i.blocking for i in issues)


def test_missing_optional_column_is_fine_and_required_error_names_closest(tmp_path, cfg):
    def dirty(wb):
        ws = wb["Portfolio"]
        hdr = [c.value for c in ws[1]]
        ws.delete_cols(hdr.index("Headcount") + 1)
    path = _edit(make_workbook(tmp_path, [position()], []), dirty)
    snap, _, issues = _issues(path, cfg)
    assert snap.positions[0].headcount is None and issues == []

    def worse(wb):
        ws = wb["Portfolio"]
        hdr = [c.value for c in ws[1]]
        ws.cell(row=1, column=hdr.index("Prior Mark ($M)") + 1).value = "Prior Mrk (USD)"   # 4 edits: too far to resolve
    path = _edit(path, worse)
    with pytest.raises(IngestError, match=r"Prior Mark \(\$M\).*closest header seen: 'Prior Mrk \(USD\)'"):
        read_workbook(path, cfg)


# ---------------------------------------------------------------- structure (§2.4) and sheets (§2.1)

def test_header_on_row_three_under_a_merged_title_loads_with_x917(tmp_path, cfg):
    def dirty(wb):
        for ws in (wb["Portfolio"], _act(wb)):
            ws.insert_rows(1, amount=2)
            ws["A1"] = "HC Portfolio — internal, Q3 2026"
            ws.merge_cells("A1:E1")
    path = _edit(make_workbook(tmp_path, [position()], [event(value=150.0, ownership_after=0.1)]), dirty)
    snap, feed, issues = _issues(path, cfg)
    assert snap.positions[0].row_index == 4 and feed.events[0].row_index == 4
    x917 = _by_rule(issues, "X-917")
    assert len(x917) == 2 and all("row 3" in i.message and not i.blocking for i in x917)
    assert not any(i.blocking for i in issues)


def test_blank_rows_total_row_and_trailing_note_are_skipped_with_x917(tmp_path, cfg):
    def dirty(wb):
        ws = wb["Portfolio"]
        ws.insert_rows(3)                       # blank row between the two positions
        ws.append(["Total", None, None, None, None, None, None, 200.0])
        ws.append(["Source: CFO pack, unaudited."])
    path = _edit(make_workbook(tmp_path, [position(), position(company="Beta")], []), dirty)
    snap, _, issues = _issues(path, cfg)
    assert [p.company for p in snap.positions] == ["Alpha", "Beta"]
    x917 = _by_rule(issues, "X-917")
    assert len(x917) == 3 and all(not i.blocking for i in x917)
    texts = " | ".join(i.message for i in x917)
    assert "1 blank row" in texts and "totals row" in texts and "only a note" in texts
    assert not any(i.blocking for i in issues)


def test_header_and_zero_activity_rows_is_an_empty_quarter(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [])
    snap, feed, issues = _issues(path, cfg)
    assert feed.events == () and issues == []


def test_zero_portfolio_rows_is_an_ingest_error(tmp_path, cfg):
    path = make_workbook(tmp_path, [], [])
    with pytest.raises(IngestError, match="no portfolio rows"):
        read_workbook(path, cfg)


@pytest.mark.parametrize("sheet,label", [
    ("Q4-2026 Activity", "Q4 2026"), ("Q4'26 Activity", "Q4 2026"), ("4Q26 Activity", "Q4 2026"),
    ("Activity Q4 2026", "Q4 2026"), ("2026 Q4 Activity", "Q4 2026"), ("Q4 2026 Events", "Q4 2026"),
    ("Activity", "Q3 2026"), ("Events", "Q3 2026"), ("Q4 Activity", "Q3 2026"), ("Activity Log", "Q3 2026"),
])
def test_relaxed_activity_sheet_names_load_with_x919(tmp_path, cfg, sheet, label):
    path = make_workbook(tmp_path, [position()], [event(value=150.0, ownership_after=0.1)], activity_sheet=sheet)
    _, feed, issues = _issues(path, cfg)
    assert feed.sheet_name == sheet and feed.quarter_label == label and len(feed.events) == 1
    x919 = _by_rule(issues, "X-919")
    assert len(x919) == 1 and not x919[0].blocking and sheet in x919[0].message
    assert not any(i.blocking for i in issues)


def test_two_activity_like_sheets_is_a_blocking_x914(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [], activity_sheet="Activity")
    _edit(path, lambda wb: wb.create_sheet("Events"))
    _, feed, issues = _issues(path, cfg)
    x914 = _by_rule(issues, "X-914")
    assert len(x914) == 1 and x914[0].blocking and "Activity" in x914[0].message and "Events" in x914[0].message


def test_portfolio_sheet_found_by_fallback_name(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [])
    _edit(path, lambda wb: setattr(wb["Portfolio"], "title", "Holdings"))
    snap, _, issues = _issues(path, cfg)
    assert snap.sheet_name == "Holdings" and [i.rule_id for i in issues] == ["X-919"]


def test_no_portfolio_sheet_by_any_rule_is_an_ingest_error(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [])
    _edit(path, lambda wb: setattr(wb["Portfolio"], "title", "Sheet1"))
    with pytest.raises(IngestError, match="portfolio sheet 'Portfolio' not found"):
        read_workbook(path, cfg)


# ---------------------------------------------------------------- hard failures (§2.5)

def test_csv_renamed_xlsx_is_not_a_workbook(tmp_path, cfg):
    path = tmp_path / "book.xlsx"
    path.write_text("Company,Sector,Fund\nAlpha,SaaS,Fund I\n")
    with pytest.raises(IngestError) as ex:
        read_workbook(path, cfg)
    assert "is not a workbook" in str(ex.value) and "CSV" in str(ex.value) and "Traceback" not in str(ex.value)


def test_empty_file_is_not_a_workbook(tmp_path, cfg):
    path = tmp_path / "empty.xlsx"
    path.write_bytes(b"")
    with pytest.raises(IngestError, match="is not a workbook.*empty file"):
        read_workbook(path, cfg)


# ---------------------------------------------------------------- currency (X-920) and terminal companies (X-907)

def test_currency_in_notes_detail_or_value_cell_blocks_with_x920(tmp_path, cfg):
    from datetime import date
    path = make_workbook(tmp_path, [position()], [
        event(value=150.0, ownership_after=0.1, notes="Round denominated in EUR"),
        event(value=150.0, ownership_after=0.1, detail="£12m Series B", date=date(2026, 8, 16)),
        event(value="€150M", ownership_after=0.1, date=date(2026, 8, 17)),
        event(value=150.0, ownership_after=0.1, notes="$28.7M round, USD, cadence of audits", date=date(2026, 8, 18)),
    ])
    _, _, issues = _issues(path, cfg)
    x920 = sorted(_by_rule(issues, "X-920"), key=lambda i: i.row_index)
    assert [i.row_index for i in x920] == [2, 3, 4] and all(i.blocking for i in x920)
    assert "EUR" in x920[0].message and "£" in x920[1].message and "€" in x920[2].message


def test_x907_allows_distribution_on_acquired_and_note_repaid_on_shut_down(tmp_path, cfg):
    from datetime import date
    path = make_workbook(tmp_path, [
        position(company="Sold", status="Acquired", prior_mark=0.0, ownership=0.0, realized=20.0),
        position(company="Dead", status="Shut Down", prior_mark=0.0, ownership=0.0, realized=0.0),
    ], [
        event(EventType.DISTRIBUTION, company="Sold", proceeds=1.0),
        event(EventType.DISTRIBUTION, company="Dead", proceeds=0.2, date=date(2026, 8, 16)),
        event(EventType.NOTE_REPAID, company="Dead", proceeds=0.5, date=date(2026, 8, 17)),
        event(EventType.NOTE_REPAID, company="Sold", proceeds=0.5, date=date(2026, 8, 18)),   # not allowed on Acquired
        event(EventType.PRICED_ROUND, company="Sold", value=1.0, ownership_after=0.1, date=date(2026, 8, 19)),
    ])
    _, _, issues = _issues(path, cfg)
    x907 = _by_rule(issues, "X-907")
    assert [i.row_index for i in x907] == [5, 6] and all(i.blocking for i in x907)


def test_new_investment_for_unknown_company_is_x918_review_not_x901(tmp_path, cfg):
    path = make_workbook(tmp_path, [position()], [
        event(EventType.NEW_INVESTMENT, company="Ghost", value=40.0, hc_investment=2.0, ownership_after=0.05,
              notes="First check from Fund III."),
    ])
    _, _, issues = _issues(path, cfg)
    assert not _by_rule(issues, "X-901")
    x918 = _by_rule(issues, "X-918")
    assert len(x918) == 1 and not x918[0].blocking and x918[0].severity.value == "REVIEW"
    assert "Ghost" in x918[0].message and "'Fund III'" in x918[0].message
