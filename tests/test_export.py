"""Export layer: workbook shape, CSVs, the E-02 re-ingest gate, and the static report."""
from __future__ import annotations

import csv
import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest
from conftest import next_quarter_cfg
import yaml

from hc_valuation.export import (
    next_quarter_label, write_csvs, write_next_quarter_workbook, write_static_report, write_workbook,
)
from hc_valuation.export.static_report import inline_bundle, run_json_script
from hc_valuation.ingest.reader import read_workbook
from hc_valuation.ingest.schema import ACTIVITY_COLUMNS, PORTFOLIO_COLUMNS
from hc_valuation.ingest.validate import validate
from hc_valuation.pipeline import RunPaths, execute, load_prior_open_items

GENERATED = datetime(2026, 10, 1, 9, 0, tzinfo=timezone.utc)


@pytest.fixture(scope="module")
def result():
    return execute(RunPaths.default(), generated_at=GENERATED, adjudicate=False)


# ------------------------------------------------------------------ workbook + csv

def test_workbook_sheets_and_row_counts(result, tmp_path: Path):
    run = result.run
    p = write_workbook(run, tmp_path / "val.xlsx")
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Summary", "Marks", "Exceptions", "Audit Trail", "Fund Rollup", "Open Items",
                             "Validation", "Alternatives"]
    assert wb["Marks"].max_row == 1 + len(run.companies) == 101
    assert wb["Exceptions"].max_row == 1 + sum(len(c.flags) for c in run.companies)
    assert wb["Audit Trail"].max_row == 1 + sum(len(c.steps) for c in run.companies)
    assert wb["Fund Rollup"].max_row == 1 + len(run.rollups)
    assert wb["Open Items"].max_row == 1 + len(run.open_items)
    assert wb["Alternatives"].max_row == 1 + sum(len(c.alternative_marks) for c in run.companies)
    for name in wb.sheetnames:
        ws = wb[name]
        assert ws.freeze_panes == "A2"
        assert ws["A1"].font.b
    marks = wb["Marks"]
    headers = [c.value for c in marks[1]]
    assert marks.cell(row=2, column=headers.index("Prior ($M)") + 1).number_format == "0.00"
    assert marks.cell(row=2, column=headers.index("Ownership After") + 1).number_format == "0.0%"
    # the only timestamp is the manifest's generated_at
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Generated at"] == GENERATED.isoformat()
    assert summary["Proposed NAV ($M)"] == pytest.approx(1183.9, abs=0.05)


def test_workbook_is_deterministic(result, tmp_path: Path):
    a = write_workbook(result.run, tmp_path / "a.xlsx")
    b = write_workbook(result.run, tmp_path / "b.xlsx")
    rows = lambda p: [[c for c in r] for ws in openpyxl.load_workbook(p).worksheets for r in ws.iter_rows(values_only=True)]  # noqa: E731
    assert rows(a) == rows(b)


def test_csvs(result, tmp_path: Path):
    from hc_valuation.api.sources import build_sources
    files = write_csvs(result.run, tmp_path, build_sources(result))
    assert {f.name for f in files} == {"marks.csv", "exceptions.csv", "audit_trail.csv", "open_items.csv", "alternatives.csv"}
    with (tmp_path / "marks.csv").open(newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 100
    dray = next(r for r in rows if r["Company"] == "Drayvenn")
    assert dray["FV Level"] == "1" and dray["Disposition"] == "BLOCK"
    with (tmp_path / "audit_trail.csv").open(newline="") as fh:
        audit = list(csv.DictReader(fh))
    assert len(audit) == sum(len(c.steps) for c in result.run.companies)
    # the workpaper chain: every step names its Portfolio row, and every input read from a cell cites it
    tarn = next(r for r in audit if r["Company"] == "Tarnwick Aerospace" and r["Rule"] == "M-012")
    assert tarn["Portfolio Row"] == "21"
    assert "post_money='Q3 2026 Activity'!E14" in tarn["Input Cells"] and "prior_post_money=Portfolio!H21" in tarn["Input Cells"]
    carry = next(r for r in audit if r["Company"] == "Beltrix" and r["Rule"] == "M-000")
    assert carry["Portfolio Row"] and "prior_mark=Portfolio!" in carry["Input Cells"]


# ------------------------------------------------------------------ E-02 snapshot gate

def test_next_quarter_label_is_parsed_not_hardcoded():
    assert next_quarter_label("Q3 2026") == "Q4 2026"
    assert next_quarter_label("Q4 2026") == "Q1 2027"
    assert next_quarter_label("Q1 2031") == "Q2 2031"
    with pytest.raises(ValueError):
        next_quarter_label("FY 2026")


def test_snapshot_reingests_with_zero_blocking_issues(result, tmp_path: Path):
    """The phase-09 gate: the emitted workbook is a valid input for the next quarter."""
    run, cfg = result.run, result.config
    out = write_next_quarter_workbook(run, result.paths.workbook, tmp_path / "next.xlsx", cfg)
    # The emitted book is next quarter's input, so it is read under next quarter's policy — the
    # object `hc-valuation next-policy` writes. Under this quarter's policy X-922 refuses it.
    assert any(i.rule_id == "X-922" and i.blocking for i in validate(*read_workbook(out, cfg), cfg))
    ncfg = next_quarter_cfg(cfg)
    snapshot, feed = read_workbook(out, ncfg)
    # Without the sidecar the deliberate departures (note at cost, pending deal) BLOCK on X-904 —
    # that is correct: an unexplained mark that is not ownership × last round must not slide through.
    unexplained = [i for i in validate(snapshot, feed, ncfg) if i.blocking]
    assert {i.company for i in unexplained} == {"Gryphonel", "Duskfern"}
    # With the sidecar the pipeline emitted beside the workbook, the same rows are non-blocking REVIEWs.
    from hc_valuation.pipeline import load_mark_basis
    issues = validate(snapshot, feed, ncfg, explained_departures=load_mark_basis(tmp_path / "open_items_carry.yaml"))
    assert len(snapshot.positions) == 100
    assert [i for i in issues if i.blocking] == []
    assert {i.company for i in issues if i.rule_id == "X-904"} == {"Gryphonel", "Duskfern"}
    # The last-round print is written as-is — never a fabricated "carrying basis" post-money.
    assert snapshot.by_company()["Gryphonel"].latest_post_money == pytest.approx(97.0)
    assert snapshot.by_company()["Duskfern"].latest_post_money == pytest.approx(73.8)
    assert feed.sheet_name == f"{next_quarter_label(cfg.quarter.label)} Activity"
    assert feed.events == ()

    wb = openpyxl.load_workbook(out)
    assert "Field Definitions" in wb.sheetnames and "Open Items" in wb.sheetnames
    ws = wb["Portfolio"]
    assert [c.value for c in ws[1]] == list(PORTFOLIO_COLUMNS)
    assert [c.value for c in wb[feed.sheet_name][1]] == list(ACTIVITY_COLUMNS)
    headers = {c.value: i for i, c in enumerate(ws[1])}
    assert re.match(r"^=\(K\d+\+L\d+\)/I\d+$", ws.cell(row=2, column=headers["MOIC (x)"] + 1).value)
    assert ws.cell(row=2, column=headers["Runway (mo)"] + 1).value.startswith("=IF(")

    by = snapshot.by_company()
    booked = run.by_company()
    for name, c in booked.items():
        p = by[name]
        assert p.prior_mark == pytest.approx(c.booked_mark if c.status_after.value == "Active" else 0.0, abs=1e-6)
        assert p.ownership == pytest.approx(c.ownership_after)
        assert p.invested == pytest.approx(c.invested_after)
        assert p.realized == pytest.approx(c.realized_cumulative)
        assert p.status == c.status_after
    assert by["Cindral"].status.value == "Acquired" and by["Cindral"].arr is None
    assert by["Drayvenn"].latest_round.isoformat() == "2026-09-20"     # IPO date, not the source round
    assert by["Fernwave"].latest_round.isoformat() == "2026-08-04"     # priced round date

    carry = tmp_path / "open_items_carry.yaml"
    assert carry.exists()
    items = load_prior_open_items(carry)
    assert {(i.company, i.kind.value) for i in items} == {(o.company, o.kind.value) for o in run.open_items}
    assert yaml.safe_load(carry.read_text())["quarter"] == cfg.quarter.label


# ------------------------------------------------------------------ static report

def test_static_report_fallback_is_self_contained(result, tmp_path: Path):
    p = write_static_report(result.run, tmp_path / "report.html", tmp_path / "no-such-static")
    html = p.read_text()
    assert "window.__HC_RUN__" in html
    assert "http" not in html.lower().replace("<!doctype html>", "")
    assert "Drayvenn" in html and "Blocked" in html
    for c in result.run.companies:
        assert c.company in html


def test_static_report_inlines_bundle(result, tmp_path: Path):
    static = tmp_path / "static"
    (static / "assets").mkdir(parents=True)
    (static / "assets" / "app.js").write_text("console.log('</script>');")
    (static / "assets" / "app.css").write_text("body{margin:0}")
    (static / "index.html").write_text(
        '<!doctype html><html><head><link rel="stylesheet" href="/assets/app.css">'
        '<link rel="modulepreload" href="/assets/app.js"></head>'
        '<body><div id="root"></div><script type="module" crossorigin src="/assets/app.js"></script></body></html>')
    p = write_static_report(result.run, tmp_path / "report.html", static)
    html = p.read_text()
    assert "window.__HC_RUN__" in html
    assert "body{margin:0}" in html and "console.log" in html
    assert 'src="/assets' not in html and "modulepreload" not in html
    assert html.index("window.__HC_RUN__") < html.index("console.log")
    assert "http" not in html.lower().replace("<!doctype html>", "")


def test_run_json_script_escapes_script_close():
    class Fake:
        def model_dump_json(self):
            return '{"rationale":"</script><script>alert(1)</script>"}'
    s = run_json_script(Fake())  # type: ignore[arg-type]
    assert "</script><script>" not in s.split("window.__HC_RUN__ = ", 1)[1].rsplit(";</script>", 1)[0]
    assert s.count("</script>") == 1
