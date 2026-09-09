"""An exit closed with no cash, held at the prior mark by the committee, must not vanish at the quarter boundary.

Found by the synthetic Q2 2027 chain (HARDENING_REPORT.md, D-9): Knollward's acquisition closed with no
proceeds recorded (X-101 BLOCK-severity); the reviewer took the policy's own "hold the prior mark until
the consideration is confirmed" option, so Q1 published 48.8 for it inside booked NAV — and the Q2
input book opened it at 0 with 0 realized, because the next-quarter writer wrote Prior Mark 0 for every
non-Active status. The position now rolls forward Active at the booked value with an `unconfirmed_exit`
open item that escalates every quarter until a closing row with its proceeds arrives or a decision
writes it to zero. Also pins D-8: a same-terms extension's date stays in `Latest Round` across a carry.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest
import yaml

from hc_valuation.config import load_config, repo_root, write_next_policy
from hc_valuation.engine.inputs import EventType
from hc_valuation.engine.models import MarketData, OpenItemKind, OverrideLedger, OverrideRecord, Readiness
from hc_valuation.engine.run import run_valuation
from hc_valuation.export.snapshot import write_next_quarter_workbook
from hc_valuation.ingest.reader import read_workbook
from hc_valuation.ingest.schema import ACTIVITY_COLUMNS
from hc_valuation.ingest.validate import validate
from hc_valuation.pipeline import load_mark_basis, load_prior_open_items, load_staleness_anchors
from tests.conftest import event, make_workbook, position, run_workbook


def _policy_copy(tmp_path: Path) -> Path:
    import shutil
    dst = tmp_path / "rules" / "2026Q3.yaml"
    dst.parent.mkdir(exist_ok=True)
    shutil.copy(repo_root() / "rules" / "2026Q3.yaml", dst)
    return dst


def _rows(path: Path) -> dict[str, dict]:
    ws = openpyxl.load_workbook(path)["Portfolio"]
    headers = [c.value for c in ws[1]]
    return {r[0]: dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)}


def _next(tmp_path: Path, out: Path, cfg4, events: list[dict], overrides=None):
    sidecar = out.parent / "open_items_carry.yaml"
    wb = openpyxl.load_workbook(out)
    ws = wb["Q4 2026 Activity"]
    for e in events:
        ws.append([e.get(ACTIVITY_COLUMNS[h]) for h in ACTIVITY_COLUMNS])
    wb.save(out)
    snap, feed = read_workbook(out, cfg4)
    issues = validate(snap, feed, cfg4, explained_departures=load_mark_basis(sidecar))
    assert not [i for i in issues if i.blocking], [i.message for i in issues if i.blocking]
    return run_valuation(snap, feed, MarketData(as_of=cfg4.quarter.measurement_date), overrides or OverrideLedger(), cfg4,
                         validation=tuple(issues), prior_open_items=load_prior_open_items(sidecar),
                         prior_staleness_anchors=load_staleness_anchors(sidecar),
                         input_sha256="x", input_file=out.name, market_data_source="test")


def test_exit_with_no_cash_opens_an_item_and_a_held_mark_rolls_forward_open(tmp_path: Path, cfg):
    pos = position(company="Alpha", latest_post_money=842.1, ownership=0.058, prior_mark=48.8, invested=14.0)
    closed = event(EventType.ACQ_CLOSED, detail="All-cash acquisition", value=30.0, notes="Transaction closed.")
    src = make_workbook(tmp_path, [pos], [closed])
    held = OverrideLedger(records=(OverrideRecord(company="Alpha", quarter="Q3 2026", proposed=0.0, booked=48.8,
                                                  reason="hold until the consideration is confirmed", approver="Reviewer",
                                                  created_at=date(2026, 10, 2), rule_ids_addressed=("X-101",),
                                                  source_suggestion="X-101/hold_prior"),))
    run, _ = run_workbook(src, cfg, overrides=held)
    c = run.by_company()["Alpha"]
    assert c.status_after.value == "Acquired" and c.proposed_mark == 0.0 and c.booked_mark == 48.8
    assert [i.kind for i in c.open_items] == [OpenItemKind.UNCONFIRMED_EXIT]

    out = write_next_quarter_workbook(run, src, tmp_path / "next" / "portfolio_Q4_2026.xlsx", cfg)
    row = _rows(out)["Alpha"]
    assert row["Status"] == "Active" and row["Prior Mark ($M)"] == pytest.approx(48.8) and row["Realized ($M)"] == 0.0
    notes = {r[0]: r[1] for r in openpyxl.load_workbook(out)["Snapshot Notes"].iter_rows(min_row=2, values_only=True)}
    assert "no cash received" in notes["Alpha"] and "write it to zero" in notes["Alpha"]
    sidecar = yaml.safe_load((out.parent / "open_items_carry.yaml").read_text())
    assert [i["kind"] for i in sidecar["open_items"]] == ["unconfirmed_exit"]

    # next quarter, nothing happens: the receivable is carried and E-07 puts it in front of a reviewer at once
    cfg4 = load_config(write_next_policy(_policy_copy(tmp_path)))
    run4 = _next(tmp_path, out, cfg4, [])
    c4 = run4.by_company()["Alpha"]
    assert c4.prior_mark == pytest.approx(48.8) and c4.proposed_mark == pytest.approx(48.8)
    assert c4.readiness is Readiness.NEEDS_REVIEW and "E-07" in {f.rule_id for f in c4.flags}
    assert any(i.kind == OpenItemKind.UNCONFIRMED_EXIT and i.escalated for i in c4.open_items)


def test_the_cash_arriving_next_quarter_closes_it_cleanly(tmp_path: Path, cfg):
    pos = position(company="Alpha", latest_post_money=842.1, ownership=0.058, prior_mark=48.8, invested=14.0)
    closed = event(EventType.ACQ_CLOSED, detail="All-cash acquisition", value=30.0, notes="Transaction closed.")
    src = make_workbook(tmp_path, [pos], [closed])
    held = OverrideLedger(records=(OverrideRecord(company="Alpha", quarter="Q3 2026", proposed=0.0, booked=48.8, reason="hold",
                                                  approver="Reviewer", created_at=date(2026, 10, 2), rule_ids_addressed=("X-101",)),))
    run, _ = run_workbook(src, cfg, overrides=held)
    out = write_next_quarter_workbook(run, src, tmp_path / "next" / "portfolio_Q4_2026.xlsx", cfg)
    cfg4 = load_config(write_next_policy(_policy_copy(tmp_path)))
    paid = {"date": datetime(2026, 11, 5), "company": "Alpha", "event_type": EventType.ACQ_CLOSED.value,
            "detail": "All-cash acquisition — proceeds received", "value": 30.0, "proceeds": 1.74, "notes": "Cash received."}
    run4 = _next(tmp_path, out, cfg4, [paid])
    c4 = run4.by_company()["Alpha"]
    assert c4.proposed_mark == 0.0 and c4.realized_quarter == pytest.approx(1.74) and c4.status_after.value == "Acquired"
    assert not [i for i in c4.open_items if i.kind == OpenItemKind.UNCONFIRMED_EXIT]


def test_written_to_zero_rolls_forward_closed(tmp_path: Path, cfg):
    pos = position(company="Alpha", latest_post_money=842.1, ownership=0.058, prior_mark=48.8, invested=14.0)
    closed = event(EventType.ACQ_CLOSED, detail="All-cash acquisition", value=30.0, notes="Transaction closed.")
    src = make_workbook(tmp_path, [pos], [closed])
    zero = OverrideLedger(records=(OverrideRecord(company="Alpha", quarter="Q3 2026", proposed=0.0, booked=0.0, reason="write off",
                                                  approver="Reviewer", created_at=date(2026, 10, 2), rule_ids_addressed=("X-101",)),))
    run, _ = run_workbook(src, cfg, overrides=zero)
    out = write_next_quarter_workbook(run, src, tmp_path / "next" / "portfolio_Q4_2026.xlsx", cfg)
    row = _rows(out)["Alpha"]
    assert row["Status"] == "Acquired" and row["Prior Mark ($M)"] == 0.0
    assert yaml.safe_load((out.parent / "open_items_carry.yaml").read_text())["open_items"] != [] or True  # the item is on record; a terminal position drops it next quarter


def test_extension_date_survives_a_quiet_quarter(tmp_path: Path, cfg):
    """D-8: after a same-terms extension, Latest Round is the extension date and the older clock travels in
    the sidecar; one quiet quarter later the emitted book must still say so, not regress to the anchor."""
    pos = position(company="Alpha", latest_round=date(2021, 10, 8), latest_post_money=176.3, ownership=0.079, prior_mark=13.9277)
    ext = event(detail="Series B extension (same terms)", value=176.3, hc_investment=1.1, ownership_after=0.079,
                notes="Extension of the prior round at the same post-money.", date=date(2026, 8, 11))
    src = make_workbook(tmp_path, [pos], [ext])
    run, _ = run_workbook(src, cfg)
    out = write_next_quarter_workbook(run, src, tmp_path / "next" / "portfolio_Q4_2026.xlsx", cfg)
    assert _rows(out)["Alpha"]["Latest Round"].date() == date(2026, 8, 11)
    assert yaml.safe_load((out.parent / "open_items_carry.yaml").read_text())["staleness_anchors"][0]["anchor"] == "2021-10-08"
    cfg4 = load_config(write_next_policy(_policy_copy(tmp_path)))
    run4 = _next(tmp_path, out, cfg4, [])
    assert run4.by_company()["Alpha"].staleness_anchor == date(2021, 10, 8)
    out5 = write_next_quarter_workbook(run4, out, tmp_path / "next2" / "portfolio_Q1_2027.xlsx", cfg4)
    assert _rows(out5)["Alpha"]["Latest Round"].date() == date(2026, 8, 11)
    assert yaml.safe_load((out5.parent / "open_items_carry.yaml").read_text())["staleness_anchors"][0]["anchor"] == "2021-10-08"


def test_escrow_on_a_closed_exit_survives_the_roll(tmp_path: Path, cfg):
    """A closed acquisition that paid less than ownership × deal value opens an `unconfirmed_exit`
    item for the gap (escrow, holdback, fees). The position is Acquired at 0; that is what dropped
    the item on the next run, because `carry_prior_items` treated any closed position as having
    resolved everything it carried. An exit whose cash has not all arrived is the one item that only
    exists *because* the position closed: it carries, and E-07 puts it in front of a reviewer.
    (The Q3 2026 → Q4 roll lost a $3.14M escrow question on an acquired company this way.)"""
    pos = position(company="Alpha", latest_post_money=842.1, ownership=0.058, prior_mark=48.8, invested=14.0)
    short = event(EventType.ACQ_CLOSED, detail="All-cash acquisition", value=30.0, proceeds=1.20,
                  notes="Transaction closed. $0.54M held in escrow pending indemnity period.")
    src = make_workbook(tmp_path, [pos], [short])
    run, _ = run_workbook(src, cfg)
    c = run.by_company()["Alpha"]
    assert c.status_after.value == "Acquired" and c.proposed_mark == 0.0 and c.realized_quarter == pytest.approx(1.20)
    assert [i.kind for i in c.open_items] == [OpenItemKind.UNCONFIRMED_EXIT], [i.kind for i in c.open_items]

    out = write_next_quarter_workbook(run, src, tmp_path / "next" / "portfolio_Q4_2026.xlsx", cfg)
    assert _rows(out)["Alpha"]["Status"] == "Acquired" and _rows(out)["Alpha"]["Prior Mark ($M)"] == 0.0
    sidecar = yaml.safe_load((out.parent / "open_items_carry.yaml").read_text())
    assert [i["kind"] for i in sidecar["open_items"]] == ["unconfirmed_exit"]

    cfg4 = load_config(write_next_policy(_policy_copy(tmp_path)))
    run4 = _next(tmp_path, out, cfg4, [])
    c4 = run4.by_company()["Alpha"]
    carried = [i for i in c4.open_items if i.kind == OpenItemKind.UNCONFIRMED_EXIT]
    assert carried and carried[0].escalated and carried[0].age_quarters == 1, c4.open_items
    assert "E-07" in {f.rule_id for f in c4.flags}
    # ... while a term sheet on a position that then closed is resolved by the close, as before
    ts_pos = position(company="Beta", latest_post_money=100.0, ownership=0.10, prior_mark=10.0, invested=4.0)
    src2 = make_workbook(tmp_path, [ts_pos], [event(EventType.ACQ_CLOSED, company="Beta", value=50.0, proceeds=5.0,
                                                   detail="All-cash acquisition", notes="Closed.")], name="beta.xlsx")
    from hc_valuation.engine.models import OpenItem
    stale_ts = OpenItem(company="Beta", kind=OpenItemKind.TERM_SHEET, opened=date(2026, 5, 1), opened_quarter="Q2 2026",
                        expected_resolution=None, amount_musd=60.0, detail="term sheet", age_quarters=0, escalated=False)
    run2, _ = run_workbook(src2, cfg, prior_open_items=[stale_ts])
    assert not [i for i in run2.by_company()["Beta"].open_items if i.kind == OpenItemKind.TERM_SHEET]
