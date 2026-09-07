"""Choosing the workbook from the dashboard (workbooks.py, GET /api/workbooks, POST /api/workbook).

A switch must carry the quarter's policy and the book's own ledger with it, must produce a
distinct run (the id derives from the inputs), and must never let a test quarter's decisions
reach the real book's ledger.
"""
from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from hc_valuation.api.app import create_app
from hc_valuation.config import repo_root, write_next_policy
from hc_valuation.pipeline import RunPaths
from hc_valuation.workbooks import (QUARTERS_DIR, SYNTHETIC_SHEET, discover, ledger_dir_for, policy_for, profile_for,
                                    provider_for, quarter_of)

ROOT = repo_root()
REAL = ROOT / "data" / "HC_Mock_Portfolio_Data.xlsx"


def _synthetic_copy(root: Path, rel: str, *, activity: str = "Q4 2026 Activity", marker: bool = True) -> Path:
    """The real workbook re-labelled as another quarter, with an empty activity tab and the
    synthetic marker sheet, placed under the scratch root."""
    dst = root / rel
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(REAL)
    act = next(ws for ws in wb.worksheets if ws.title.endswith("Activity"))
    act.delete_rows(2, act.max_row)
    act.title = activity
    if marker:
        wb.create_sheet(SYNTHETIC_SHEET).append(["This workbook is invented test data."])
    wb.save(dst)
    return dst


@pytest.fixture()
def root(tmp_path: Path) -> Path:
    """A scratch root with the base policy only: the repo may carry later quarters' policy files
    (the synthetic chain writes rules/2026Q4.yaml and on), and these tests build that state
    themselves with write_next_policy where they need it."""
    shutil.copytree(ROOT / "rules", tmp_path / "rules")
    for p in (tmp_path / "rules").glob("*.yaml"):
        if p.name not in ("2026Q3.yaml", "comps_baskets.yaml", "rationale.yaml"):
            p.unlink()
    (tmp_path / "data").mkdir()
    shutil.copy(REAL, tmp_path / "data" / REAL.name)
    shutil.copytree(ROOT / "data" / "mock_responses", tmp_path / "data" / "mock_responses")
    shutil.copy(ROOT / "data" / "overrides.yaml", tmp_path / "data" / "overrides.yaml")
    return tmp_path


# ------------------------------------------------------------------ pieces

def test_quarter_and_marker_are_read_from_the_sheet_names(root: Path):
    assert quarter_of(REAL) == ("Q3 2026", False)
    wb = _synthetic_copy(root, "data/quarters/chain-a/Q4_2026.xlsx")
    assert quarter_of(wb) == ("Q4 2026", True)
    assert quarter_of(root / "nope.xlsx") == (None, False)


def test_policy_and_ledger_follow_the_workbook(root: Path):
    assert policy_for(root, "Q3 2026") == root / "rules" / "2026Q3.yaml"
    assert policy_for(root, "Q4 2026") is None
    write_next_policy(root / "rules" / "2026Q3.yaml")
    assert policy_for(root, "Q4 2026") == root / "rules" / "2026Q4.yaml"
    # a real book records into the committee's ledger wherever it sits; a synthetic one gets its chain's
    assert ledger_dir_for(root, root / "data" / REAL.name) == root / "data"
    assert ledger_dir_for(root, root / QUARTERS_DIR / "2026Q4" / "portfolio_Q4_2026.xlsx") == root / "data"
    assert ledger_dir_for(root, root / QUARTERS_DIR / "chain-a" / "Q4_2026.xlsx", synthetic=True) == root / QUARTERS_DIR / "chain-a" / "ledger"
    assert ledger_dir_for(root, root / QUARTERS_DIR / "chain-a" / "2027Q1" / "book.xlsx", synthetic=True) == root / QUARTERS_DIR / "chain-a" / "ledger"


def test_run_paths_default_policy_follows_the_workbook_quarter(root: Path):
    """`hc-valuation run --input <a Q4 book>` with no --policy must not run it under the Q3 policy."""
    wb = _synthetic_copy(root, "data/quarters/2026Q4/portfolio_Q4_2026.xlsx", marker=False)
    assert RunPaths.default(root=root, workbook=wb).policy == root / "rules" / "2026Q3.yaml"   # no Q4 policy yet: the base
    write_next_policy(root / "rules" / "2026Q3.yaml")
    assert RunPaths.default(root=root, workbook=wb).policy == root / "rules" / "2026Q4.yaml"
    assert RunPaths.default(root=root).policy == root / "rules" / "2026Q3.yaml"


def test_provider_rule_is_explicit_then_synthetic_then_cache_then_fixture(root: Path, monkeypatch):
    monkeypatch.delenv("HC_MARKET_PROVIDER", raising=False)
    pol = root / "rules" / "2026Q3.yaml"
    assert provider_for(root, pol, "stub") == "stub"
    assert provider_for(root, pol) is None                      # scratch root: no cache, no synthetic file
    (root / "data" / "market_cache" / "2026-09-30").mkdir(parents=True)
    (root / "data" / "market_cache" / "2026-09-30" / "meta.json").write_text("{}")
    assert provider_for(root, pol) == "live"
    (root / "data" / "synthetic_market").mkdir()
    (root / "data" / "synthetic_market" / "2026-09-30.yaml").write_text("synthetic: true\nsectors: {}\n")
    assert provider_for(root, pol) == "synthetic"
    monkeypatch.setenv("HC_MARKET_PROVIDER", "stub")
    assert provider_for(root, pol) is None                      # the environment wins and the caller passes it through


def test_discover_lists_real_quarters_and_never_synthetic_ones_beside_the_real_book(root: Path):
    _synthetic_copy(root, "data/quarters/2026Q4/portfolio_Q4_2026.xlsx", marker=False)          # a real next quarter
    _synthetic_copy(root, "data/quarters/synthetic/Q4_2026.xlsx")                                # invented, marked by its sheet
    _synthetic_copy(root, "data/quarters/other/SYNTHETIC_portfolio_Q1_2027.xlsx", activity="Q1 2027 Activity", marker=False)   # by name
    _synthetic_copy(root, "data/quarters/2026Q4/malformed/Q4_2026_broken.xlsx", activity="Activity Q4", marker=False)
    (root / "data" / "quarters" / "2026Q4" / "~$portfolio_Q4_2026.xlsx").write_bytes(b"")
    write_next_policy(root / "rules" / "2026Q3.yaml")
    profs = discover(root, root / "data" / REAL.name)
    assert profs[0].current and profs[0].id == "data/HC_Mock_Portfolio_Data.xlsx" and not profs[0].synthetic
    assert profs[0].policy == "rules/2026Q3.yaml" and profs[0].ledger_dir == "data" and profs[0].usable
    assert [p.id for p in profs] == ["data/HC_Mock_Portfolio_Data.xlsx", "data/quarters/2026Q4/malformed/Q4_2026_broken.xlsx",
                                     "data/quarters/2026Q4/portfolio_Q4_2026.xlsx"]
    q4 = profs[2]
    assert q4.quarter == "Q4 2026" and not q4.synthetic and q4.usable and q4.policy == "rules/2026Q4.yaml" and q4.ledger_dir == "data"
    assert not profs[1].usable and "Activity" in profs[1].reason


def test_synthetic_quarters_are_offered_only_from_a_synthetic_book(root: Path, monkeypatch):
    monkeypatch.delenv("HC_INCLUDE_SYNTHETIC", raising=False)
    syn = _synthetic_copy(root, "data/quarters/synthetic/Q4_2026.xlsx")
    write_next_policy(root / "rules" / "2026Q3.yaml")
    assert all(not p.synthetic for p in discover(root, root / "data" / REAL.name))
    from_syn = discover(root, syn)
    assert from_syn[0].current and from_syn[0].synthetic and from_syn[0].ledger_dir == "data/quarters/synthetic/ledger"
    assert "data/HC_Mock_Portfolio_Data.xlsx" in [p.id for p in from_syn]   # the real book is always reachable
    monkeypatch.setenv("HC_INCLUDE_SYNTHETIC", "1")
    assert any(p.synthetic for p in discover(root, root / "data" / REAL.name))


def test_missing_policy_makes_a_profile_unusable_and_says_what_to_run(root: Path):
    wb = _synthetic_copy(root, "data/quarters/chain-a/Q4_2026.xlsx")
    p = profile_for(root, wb)
    assert not p.usable and "rules/2026Q4.yaml" in p.reason and "next-policy" in p.reason


# ------------------------------------------------------------------ the API

def test_switching_workbooks_changes_the_run_and_the_ledger_and_back(root: Path, tmp_path: Path):
    _synthetic_copy(root, "data/quarters/chain-a/Q4_2026.xlsx")   # marked synthetic: its ledger is its own, and it is served here on purpose
    write_next_policy(root / "rules" / "2026Q3.yaml")
    paths = RunPaths.default(root=root, workbook=root / "data" / REAL.name, policy=root / "rules" / "2026Q3.yaml")
    import os
    os.environ["HC_INCLUDE_SYNTHETIC"] = "1"      # the switch-and-back mechanics, exercised on the marked copy
    client = TestClient(create_app(paths, provider="stub", static_dir=tmp_path / "no-static"))
    h0 = client.get("/api/health").json()
    listing = client.get("/api/workbooks").json()
    assert listing["current"] == "data/HC_Mock_Portfolio_Data.xlsx"
    assert [w["id"] for w in listing["workbooks"]] == ["data/HC_Mock_Portfolio_Data.xlsx", "data/quarters/chain-a/Q4_2026.xlsx"]

    r = client.post("/api/workbook", json={"id": "data/quarters/chain-a/Q4_2026.xlsx"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["quarter"] == "Q4 2026" and body["run_id"] != h0["run_id"]
    assert body["ledger_dir"].endswith("data/quarters/chain-a/ledger")
    h1 = client.get("/api/health").json()
    assert h1["quarter"] == "Q4 2026" and h1["run_id"] == body["run_id"]
    assert h1["ledger"]["overrides"].endswith("data/quarters/chain-a/ledger/overrides.yaml")
    run = client.get("/api/run").json()
    assert run["manifest"]["input_file"] == "Q4_2026.xlsx" and run["manifest"]["quarter_label"] == "Q4 2026"
    assert client.get("/api/workbooks").json()["current"] == "data/quarters/chain-a/Q4_2026.xlsx"

    # a decision recorded on the test quarter lands in its ledger, not the real book's
    before = (root / "data" / "overrides.yaml").read_bytes()
    target = next(c for c in run["companies"] if c["readiness"] != "Ready")
    rules = [f["rule_id"] for f in target["flags"] if f["severity"] in ("BLOCK", "REVIEW")]
    r = client.post("/api/overrides", json={"company": target["company"], "booked": target["proposed_mark"],
                                            "reason": "test", "approver": "T", "rule_ids_addressed": rules})
    assert r.status_code == 200, r.text
    assert (root / "data" / "quarters" / "chain-a" / "ledger" / "overrides.yaml").exists()
    assert (root / "data" / "overrides.yaml").read_bytes() == before

    # and back: the same run id as before, because the same inputs
    r = client.post("/api/workbook", json={"id": "data/HC_Mock_Portfolio_Data.xlsx"})
    assert r.status_code == 200 and r.json()["run_id"] == h0["run_id"]
    assert client.get("/api/health").json()["ledger"]["overrides"].endswith("data/overrides.yaml")
    os.environ.pop("HC_INCLUDE_SYNTHETIC", None)


def test_switching_to_an_unknown_or_unusable_workbook_is_refused(root: Path, tmp_path: Path, monkeypatch):
    monkeypatch.delenv("HC_INCLUDE_SYNTHETIC", raising=False)
    _synthetic_copy(root, "data/quarters/2026Q4/portfolio_Q4_2026.xlsx", marker=False)     # real, but no 2026Q4 policy written
    syn = _synthetic_copy(root, "data/quarters/chain-a/Q4_2026.xlsx")
    paths = RunPaths.default(root=root, workbook=root / "data" / REAL.name, policy=root / "rules" / "2026Q3.yaml")
    client = TestClient(create_app(paths, provider="stub", static_dir=tmp_path / "no-static"))
    before = client.get("/api/health").json()["run_id"]
    assert client.post("/api/workbook", json={"id": "data/quarters/nope.xlsx"}).status_code == 404
    assert client.post("/api/workbook", json={"id": "data/quarters/chain-a/Q4_2026.xlsx"}).status_code == 404   # synthetic: not offered from the real book
    assert syn.exists()
    r = client.post("/api/workbook", json={"id": "data/quarters/2026Q4/portfolio_Q4_2026.xlsx"})
    assert r.status_code == 409 and "next-policy" in r.json()["detail"]["message"]
    assert client.get("/api/health").json()["run_id"] == before
