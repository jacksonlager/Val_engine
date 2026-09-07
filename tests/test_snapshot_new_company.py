"""A company that entered the book from a New Investment row must roll into next quarter's workbook.

Found by the synthetic Q4 2026 chain (HARDENING_REPORT.md, D-4): `write_next_quarter_workbook`
raised for a position the run had synthesised from activity (M-014 on a company with no Portfolio
row), which stopped `hc-valuation build` for the whole quarter. The emitted row is now built from
the run's own figures with blank operating metrics and a Snapshot Notes line saying so.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import pytest

from hc_valuation.engine.inputs import EventType
from hc_valuation.export.snapshot import write_next_quarter_workbook
from hc_valuation.ingest.reader import read_workbook
from hc_valuation.ingest.validate import validate
from tests.conftest import event, make_workbook, position, run_workbook


def test_new_company_from_activity_rolls_forward(tmp_path: Path, cfg):
    book = [position(company="Alpha")]
    entry = event(EventType.NEW_INVESTMENT, company="Newco Labs", date=date(2026, 8, 20),
                  detail="Seed — Fund III, SaaS", value=18.0, hc_investment=1.5, ownership_after=0.083,
                  notes="First cheque from Fund III into a SaaS seed round.")
    src = make_workbook(tmp_path, book, [entry])
    run, issues = run_workbook(src, cfg)
    newco = run.by_company()["Newco Labs"]
    assert newco.proposed_mark == pytest.approx(0.083 * 18.0) and {f.rule_id for f in newco.flags} >= {"X-918", "X-120"}

    out = write_next_quarter_workbook(run, src, tmp_path / "next" / "portfolio_Q4_2026.xlsx", cfg)
    wb = openpyxl.load_workbook(out)
    ws = wb["Portfolio"]
    headers = [c.value for c in ws[1]]
    rows = {r[0]: dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "Newco Labs" in rows
    row = rows["Newco Labs"]
    assert row["Fund"] == "Fund III" and row["Sector"] == "SaaS" and row["Stage"] == "Seed" and row["Status"] == "Active"
    assert row["Latest Post-Money ($M)"] == 18.0 and row["Ownership (FD %)"] == pytest.approx(0.083)
    assert row["Prior Mark ($M)"] == pytest.approx(0.083 * 18.0) and row["Invested ($M)"] == pytest.approx(1.5)
    assert row["Latest Round"].date() == date(2026, 8, 20) and row["First Investment"].date() == date(2026, 8, 20)
    assert row["ARR ($M)"] is None and row["Cash ($M)"] is None and row["Headcount"] is None
    notes = {r[0]: r[1] for r in wb["Snapshot Notes"].iter_rows(min_row=2, values_only=True)}
    assert "Newco Labs" in notes and "New Investment" in notes["Newco Labs"] and "blank" in notes["Newco Labs"]

    # the emitted book re-ingests as next quarter's input with nothing blocking
    from hc_valuation.config import load_config, write_next_policy
    cfg4 = load_config(write_next_policy(_policy_copy(tmp_path)))
    snap, feed = read_workbook(out, cfg4)
    assert "Newco Labs" in snap.by_company()
    blocking = [i for i in validate(snap, feed, cfg4) if i.blocking]
    assert blocking == []


def _policy_copy(tmp_path: Path) -> Path:
    import shutil
    from hc_valuation.config import repo_root
    dst = tmp_path / "rules" / "2026Q3.yaml"
    dst.parent.mkdir(exist_ok=True)
    shutil.copy(repo_root() / "rules" / "2026Q3.yaml", dst)
    return dst


def test_synthetic_marker_travels_with_the_emitted_book(tmp_path: Path, cfg):
    """A next-quarter input emitted from a synthetic workbook is synthetic too, and must say so
    the way the source did: the marker sheet, first. A real source emits no marker."""
    from hc_valuation.workbooks import SYNTHETIC_SHEET, quarter_of
    src = make_workbook(tmp_path, [position(company="Alpha")], [])
    wb = openpyxl.load_workbook(src)
    wb.create_sheet(SYNTHETIC_SHEET, 0).append(["This workbook is invented test data."])
    wb.save(src)
    run, _ = run_workbook(src, cfg)
    out = write_next_quarter_workbook(run, src, tmp_path / "next" / "portfolio_Q4_2026.xlsx", cfg)
    assert openpyxl.load_workbook(out).sheetnames[0] == SYNTHETIC_SHEET and quarter_of(out) == ("Q4 2026", True)

    real = make_workbook(tmp_path, [position(company="Alpha")], [], name="real.xlsx")
    run2, _ = run_workbook(real, cfg)
    out2 = write_next_quarter_workbook(run2, real, tmp_path / "next2" / "portfolio_Q4_2026.xlsx", cfg)
    assert SYNTHETIC_SHEET not in openpyxl.load_workbook(out2).sheetnames and quarter_of(out2) == ("Q4 2026", False)
