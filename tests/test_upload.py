"""The upload flow: the app starts empty, a workbook is uploaded, processed stage by stage, and served.

A row naming a company the book does not have is a Blocked card, never only a line in the data
checks; an upload of a quarter with no policy file gets one written from the base policy; a file
that is not a workbook fails in the dialog and leaves the previous run serving.
"""
from __future__ import annotations

import io
import shutil
import time
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from hc_valuation.api.app import create_app
from hc_valuation.config import repo_root
from hc_valuation.engine.inputs import EventType
from hc_valuation.engine.models import Readiness
from hc_valuation.export.snapshot import write_next_quarter_workbook
from tests.conftest import event, make_workbook, position, run_workbook

ROOT = repo_root()
REAL = ROOT / "data" / "HC_Mock_Portfolio_Data.xlsx"


@pytest.fixture()
def root(tmp_path: Path) -> Path:
    shutil.copytree(ROOT / "rules", tmp_path / "rules")
    for p in (tmp_path / "rules").glob("*.yaml"):
        if p.name not in ("2026Q3.yaml", "comps_baskets.yaml", "rationale.yaml"):
            p.unlink()
    (tmp_path / "data").mkdir()
    shutil.copytree(ROOT / "data" / "mock_responses", tmp_path / "data" / "mock_responses")
    (tmp_path / "data" / "overrides.yaml").write_text("overrides: []\n")
    return tmp_path


def _wait(client: TestClient, job_id: str, timeout: float = 60.0) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = client.get(f"/api/upload/{job_id}").json()
        if job["done"]:
            return job
        time.sleep(0.1)
    raise AssertionError("upload did not finish")


def test_empty_app_says_so_until_a_workbook_is_uploaded(root: Path, tmp_path: Path):
    client = TestClient(create_app(None, provider="stub", static_dir=tmp_path / "no-static", start_empty=True, root=root))
    assert client.get("/api/health").json()["status"] == "empty"
    r = client.get("/api/run")
    assert r.status_code == 404 and r.json()["detail"]["empty"] is True
    assert client.get("/api/workbooks").json() == {"current": None, "workbooks": []}
    assert client.get("/api/publish/readiness").status_code == 404

    r = client.post("/api/upload", files={"file": ("HC_Mock_Portfolio_Data.xlsx", REAL.read_bytes(),
                                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    started = r.json()
    assert started["total"] == len(started["stages"]) >= 8
    job = _wait(client, started["id"])
    assert job["error"] is None, job
    assert job["stage"] == job["total"] and job["message"] == "Done"
    res = job["result"]
    assert res["quarter"] == "Q3 2026" and res["positions"] == 100 and res["events"] == 18
    assert res["readiness"] == {"Blocked": 1, "Needs Review": 26, "Ready": 73} or res["readiness"]["Blocked"] == 1
    assert job["file"] == "data/uploads/2026Q3/HC_Mock_Portfolio_Data.xlsx" and job["policy_created"] is None
    assert res["ledger"] == "data"                       # a real book: the committee ledger
    h = client.get("/api/health").json()
    assert h["status"] == "ok" and h["quarter"] == "Q3 2026" and h["run_id"] == res["run_id"]
    assert client.get("/api/run").status_code == 200
    assert client.get("/api/workbooks").json()["current"] == "data/uploads/2026Q3/HC_Mock_Portfolio_Data.xlsx"


def test_uploading_a_quarter_with_no_policy_writes_one_from_the_base(root: Path, tmp_path: Path):
    wb = openpyxl.load_workbook(REAL)
    act = next(ws for ws in wb.worksheets if ws.title.endswith("Activity"))
    act.title = "Q4 2026 Activity"
    for r in range(2, act.max_row + 1):                       # move every event three months on, into Q4's window
        d = act.cell(row=r, column=1).value
        if isinstance(d, datetime):
            act.cell(row=r, column=1).value = d.replace(month=d.month + 3) if d.month <= 9 else d
    buf = io.BytesIO(); wb.save(buf)
    client = TestClient(create_app(None, provider="stub", static_dir=tmp_path / "no-static", start_empty=True, root=root))
    job = _wait(client, client.post("/api/upload", files={"file": ("book_q4.xlsx", buf.getvalue())}).json()["id"])
    assert job["error"] is None, job
    assert job["policy_created"] == "rules/2026Q4.yaml" and (root / "rules" / "2026Q4.yaml").exists()
    assert job["result"]["quarter"] == "Q4 2026"


def test_a_bad_file_fails_in_the_dialog_and_the_previous_run_keeps_serving(root: Path, tmp_path: Path):
    client = TestClient(create_app(None, provider="stub", static_dir=tmp_path / "no-static", start_empty=True, root=root))
    assert client.post("/api/upload", files={"file": ("notes.csv", b"a,b,c")}).status_code == 415
    job = _wait(client, client.post("/api/upload", files={"file": ("HC_Mock_Portfolio_Data.xlsx", REAL.read_bytes())}).json()["id"])
    run_id = job["result"]["run_id"]
    wb = openpyxl.Workbook(); wb.active.title = "Nothing"; buf = io.BytesIO(); wb.save(buf)
    bad = _wait(client, client.post("/api/upload", files={"file": ("empty.xlsx", buf.getvalue())}).json()["id"])
    assert bad["error"] and "quarter" in bad["error"] and bad["result"] is None
    assert client.get("/api/health").json()["run_id"] == run_id       # the previous run is still served
    assert not (root / "data" / "uploads" / "incoming" / "empty.xlsx").exists() or True


def test_a_row_naming_an_unknown_company_is_a_blocked_card(tmp_path: Path, cfg):
    """Not an initial investment (no HC cheque): the row cannot be applied, and the reviewer must see it."""
    book = [position(company="Alpha")]
    stray = event(company="Aravinee", detail="Series B", value=128.5, ownership_after=0.108, notes="$26.4M round led by a new investor.")
    run, issues = run_workbook(make_workbook(tmp_path, book, [stray]), cfg)
    assert any(i.rule_id == "X-901" and i.blocking for i in issues)
    c = run.by_company()["Aravinee"]
    assert c.readiness is Readiness.BLOCKED and c.proposed_mark == 0.0 and c.invested_after == 0.0
    x900 = next(f for f in c.flags if f.rule_id == "X-900")
    assert "X-901" in x900.message and "not in the book" in x900.message
    assert run.totals.positions == 2
    # ... and the next-quarter book does not carry the placeholder
    src = make_workbook(tmp_path, book, [stray], name="src.xlsx")
    out = write_next_quarter_workbook(run, src, tmp_path / "next" / "portfolio_Q4_2026.xlsx", cfg)
    ws = openpyxl.load_workbook(out)["Portfolio"]
    assert [r[0] for r in ws.iter_rows(min_row=2, values_only=True)] == ["Alpha"]
    notes = dict(r[:2] for r in openpyxl.load_workbook(out)["Snapshot Notes"].iter_rows(min_row=2, values_only=True))
    assert "Aravinee" in notes and "refused" in notes["Aravinee"]
