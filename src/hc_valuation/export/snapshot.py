"""E-02 — the next-quarter input workbook.

The booked marks become next quarter's `Prior Mark ($M)`, in the exact schema the reader
expects, so the refresh workflow is "drop the new activity into the emitted file". The
gate for this module is that the emitted file re-ingests with zero blocking issues.

The one place that gate bites is the identity the source book defines
(`Prior Mark = Ownership x Latest Post-Money`), which a probability-weighted announced
deal, a note leg carried at cost or a committee override deliberately does not satisfy.
We do NOT fake a post-money to make the identity hold — the last-round print is written
as-is, because it is the true last price and next quarter's down-round test depends on it.
Instead each departure is recorded in the `open_items_carry.yaml` sidecar under
`mark_basis`, which the pipeline hands to X-904 so the reconciliation becomes a
non-blocking REVIEW ("departs from last-round pricing; explained by ...") rather than a
BLOCK. The same rows are listed on the `Snapshot Notes` tab for the human reader.
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..config import RuleConfig
from ..engine.inputs import EventType, Position, Status
from ..engine.models import CompanyResult, ValuationRun
from ..ingest.reader import read_workbook
from ..ingest.schema import ACTIVITY_COLUMNS, PORTFOLIO_COLUMNS

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="000000")
_QUARTER_RX = re.compile(r"^\s*Q([1-4])\s+(\d{4})\s*$")

# Number formats copied from the source book so the emitted file looks like the input.
_FORMATS: dict[str, str] = {
    "First Investment": "mm/dd/yyyy", "Latest Round": "mm/dd/yyyy",
    "Latest Post-Money ($M)": r"\$#,##0.0", "Invested ($M)": r"\$#,##0.0", "Ownership (FD %)": "0.0%",
    "Prior Mark ($M)": r"\$#,##0.0", "Realized ($M)": r"\$#,##0.0", "MOIC (x)": r"0.0\x",
    "ARR ($M)": r"\$#,##0.0", "ARR Growth (YoY %)": "0%", "Gross Margin (%)": "0%",
    "Net Burn ($M/mo)": r"\$#,##0.00", "Cash ($M)": r"\$#,##0.0", "Runway (mo)": "0.0", "Headcount": "#,##0",
}
_WIDTHS: dict[str, float] = {"Company": 22, "Sector": 15, "Fund": 9, "Stage": 10, "Status": 11}
_PRICING_EVENTS = {EventType.PRICED_ROUND.value, EventType.IPO.value}
_OPERATING = ("ARR ($M)", "ARR Growth (YoY %)", "Gross Margin (%)", "Net Burn ($M/mo)", "Cash ($M)", "Headcount")


def next_quarter_label(label: str) -> str:
    """'Q3 2026' -> 'Q4 2026'; 'Q4 2026' -> 'Q1 2027'. Parsed, never hardcoded."""
    m = _QUARTER_RX.match(label)
    if not m:
        raise ValueError(f"quarter label {label!r} is not of the form 'Qn YYYY'")
    q, y = int(m.group(1)), int(m.group(2))
    return f"Q1 {y + 1}" if q == 4 else f"Q{q + 1} {y}"


def _latest_round_date(c: CompanyResult, src: Position) -> date:
    """`Latest Round` as the column defines it: the date of the most recent priced round (an IPO
    counts; a same-terms extension, M-011, is a priced round too). The staleness clock is a
    different thing — an extension does not reset it — and travels in the sidecar
    (`staleness_anchors`, see `_carried_anchors`) rather than by overloading this column."""
    dates = [s.evidence.date for s in c.steps if s.evidence and s.evidence.event_type in _PRICING_EVENTS
             and s.rule_id != "M-000"]
    if dates:
        return max(dates)
    if c.listed and c.staleness_anchor != date.min:
        return c.staleness_anchor        # a listed carry (M-041) is priced at the measurement date
    # No priced round this quarter: the column keeps the value the book already carried. Falling back
    # to the staleness anchor here rewrote a same-terms extension's date with the older clock one
    # quarter later, and lost the sidecar anchor with it (synthetic Q2 2027 chain, D-8).
    return src.latest_round if src.latest_round else c.staleness_anchor


def _carried_anchors(run: ValuationRun, sources: dict[str, Position]) -> list[dict[str, str]]:
    """Positions whose staleness anchor is older than the `Latest Round` the emitted book will
    carry — the clock the engine kept running through an extension — for the sidecar."""
    out = []
    for c in run.companies:
        src = sources.get(c.company)
        if src is None or c.staleness_anchor == date.min:
            continue
        if c.staleness_anchor < _latest_round_date(c, src):
            out.append({"company": c.company, "anchor": c.staleness_anchor.isoformat(),
                        "reason": f"same-terms extension on {_latest_round_date(c, src).isoformat()} is not price discovery; "
                                  f"the staleness clock runs from {c.staleness_anchor.isoformat()}"})
    return out


def _post_money(c: CompanyResult, tol: float) -> tuple[float, str | None]:
    """The last-round post-money, written as-is, plus a reason when the booked mark deliberately
    departs from ownership × that print (the departure is carried in the sidecar for X-904)."""
    post = c.latest_post_money
    if (c.status_after != Status.ACTIVE and not _carried_open(c)) or c.ownership_after <= 0:
        return post, None
    if abs(c.ownership_after * post - c.booked_mark) <= tol:
        return post, None
    why = []
    if c.note_at_cost:
        why.append(f"note leg ${c.note_at_cost:.2f}M carried at cost (M-060)")
    if c.override is not None:
        why.append(f"committee override booked ${c.booked_mark:.2f}M vs proposed ${c.proposed_mark:.2f}M (E-01)")
    marking_rules = [s.rule_id for s in c.steps if s.rule_id.startswith("M-") and s.rule_id not in ("M-000", "M-060", "M-080")]
    if abs(c.equity_mark - c.ownership_after * post) > tol and marking_rules:
        why.append(f"equity mark ${c.equity_mark:.2f}M set by {marking_rules[-1]}, not ownership × last round")
    reason = "; ".join(why or ["booked mark does not reconcile to ownership × last-round post-money"])
    return post, (f"Prior Mark ${c.booked_mark:.2f}M departs from ownership × last-round post-money "
                  f"({c.ownership_after:.1%} × ${post:.2f}M = ${c.ownership_after * post:.2f}M): {reason}.")


def _carried_open(c: CompanyResult) -> bool:
    """A position the quarter closed but the committee kept on the book: an exit with no cash recorded
    where the decision was to hold the prior mark. Its value is in the published NAV, so the next
    quarter must open with it rather than with a zero — the book would otherwise lose the amount
    with no cash against it (synthetic Q2 2027 chain, D-9). It rolls as Active; the open item M-020
    raised keeps it in front of a reviewer until the proceeds are recorded or it is written to zero."""
    return c.status_after != Status.ACTIVE and c.booked_mark > 0


def _portfolio_row(c: CompanyResult, src: Position, r: int, tol: float) -> tuple[list[Any], str | None]:
    active = c.status_after == Status.ACTIVE or _carried_open(c)
    post, note = _post_money(c, tol)
    if _carried_open(c):
        note = (f"Exit recorded as {c.status_after.value} in {c.steps[-1].evidence.date.isoformat() if c.steps[-1].evidence else 'the quarter'} "
                f"with no cash received; the committee held ${c.booked_mark:.2f}M rather than writing the position off, so it rolls "
                "forward as Active at that value with an 'unconfirmed exit' open item. Record the closing with its proceeds when the "
                "cash arrives, or write it to zero.") + (f" {note}" if note else "")
    values: dict[str, Any] = {
        "Company": c.company, "Sector": c.sector, "Fund": c.fund, "Stage": c.stage,
        "Status": Status.ACTIVE.value if _carried_open(c) else c.status_after.value,
        "First Investment": src.first_investment, "Latest Round": _latest_round_date(c, src),
        "Latest Post-Money ($M)": post, "Invested ($M)": c.invested_after, "Ownership (FD %)": c.ownership_after,
        "Prior Mark ($M)": c.booked_mark if active else 0.0, "Realized ($M)": c.realized_cumulative,
        "MOIC (x)": f"=(K{r}+L{r})/I{r}",
        "ARR ($M)": src.arr if active else None, "ARR Growth (YoY %)": src.arr_growth if active else None,
        "Gross Margin (%)": src.gross_margin if active else None, "Net Burn ($M/mo)": src.net_burn if active else None,
        "Cash ($M)": src.cash if active else None,
        "Runway (mo)": f'=IF(AND(ISNUMBER(Q{r}),Q{r}>0),R{r}/Q{r},"-")',
        "Headcount": src.headcount if active else None,
    }
    return [values[col] for col in PORTFOLIO_COLUMNS], note


def _style_header(ws, ncols: int, freeze: str) -> None:
    for i in range(1, ncols + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = freeze


def _entered_from_activity(c: CompanyResult, run: ValuationRun) -> Position:
    """A position the quarter created from a `New Investment` row (M-014 on a company the Portfolio
    tab did not have, X-918). The source workbook has no row to carry metrics from, so the emitted
    row is built from the run's own figures — fund, sector and stage as the row named them, the
    entry date as first investment and latest round, the entry price as the last post-money — and
    the operating metrics are left blank: they are not on file, and a blank is honest where a guess
    is not. Found by the synthetic Q4 2026 chain, where one such company stopped `build` outright."""
    entry = next((s.evidence.date for s in c.steps if s.rule_id == "M-014" and s.evidence), None) or c.staleness_anchor
    return Position(
        company=c.company, sector=c.sector, fund=c.fund, stage=c.stage, status=c.status_after,
        first_investment=entry, latest_round=entry, latest_post_money=c.latest_post_money,
        invested=c.invested_after, ownership=c.ownership_after, prior_mark=c.booked_mark, realized=c.realized_cumulative,
        row_index=0, extra={"entered_from_activity": run.manifest.quarter_label},
    )


def _write_portfolio(wb: Workbook, run: ValuationRun, sources: dict[str, Position], tol: float) -> list[tuple[str, str]]:
    ws = wb.active
    ws.title = "Portfolio"
    cols = list(PORTFOLIO_COLUMNS)
    ws.append(cols)
    notes: list[tuple[str, str]] = []
    i = 1
    for c in run.companies:
        src = sources.get(c.company)
        if src is None and c.invested_after == 0 and c.booked_mark == 0 and c.ownership_after == 0:
            # a placeholder for a refused row naming a company the book did not have: nothing to carry
            notes.append((c.company, "Named on the activity tab but not in the Portfolio tab, and the row was refused; "
                                     "no position is carried. Correct the company name, or add the row as a New Investment."))
            continue
        i += 1
        if src is None:
            src = _entered_from_activity(c, run)
            notes.append((c.company, f"Entered the book in {run.manifest.quarter_label} from a New Investment row (M-014, X-918): "
                                     f"fund, sector and stage are as that row named them, and the operating metrics (ARR, growth, "
                                     "margin, burn, cash, headcount) are blank because nothing is on file. Fill them before the next run."))
        row, note = _portfolio_row(c, src, i, tol)
        ws.append(row)
        if note:
            notes.append((c.company, note))
        for j, col in enumerate(cols, start=1):
            if col in _FORMATS:
                ws.cell(row=i, column=j).number_format = _FORMATS[col]
    for j, col in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = _WIDTHS.get(col, 13)
    _style_header(ws, len(cols), "B2")
    return notes


def _write_activity(wb: Workbook, sheet_name: str) -> None:
    ws = wb.create_sheet(sheet_name)
    cols = list(ACTIVITY_COLUMNS)
    ws.append(cols)
    for j, col in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = {"Detail": 30, "Notes": 60, "Company": 22, "Event": 21}.get(col, 13)
    _style_header(ws, len(cols), "A2")


def _copy_field_definitions(wb: Workbook, source_path: Path, run: ValuationRun, next_label: str) -> None:
    src_wb = openpyxl.load_workbook(source_path, read_only=False, data_only=True)
    name = next((s for s in src_wb.sheetnames if s.lower().startswith("field def")), None)
    ws = wb.create_sheet(name or "Field Definitions")
    if name is not None:
        src = src_wb[name]
        for row in src.iter_rows(values_only=True):
            ws.append(list(row))
        for key, dim in src.column_dimensions.items():
            ws.column_dimensions[key].width = dim.width
    else:
        ws.append(["Field", "Definition"])
    src_wb.close()
    m = run.manifest
    ws.append([None, None])
    ws.append(["SNAPSHOT", None])
    ws.append(["Generated by", f"hc-valuation engine {m.engine_version}, policy {m.policy_version}, run {m.run_id}"])
    ws.append(["Portfolio tab", f"The book as of the {m.quarter_label} close ({m.measurement_date.isoformat()}): "
                                "Prior Mark is the booked mark from that run; Ownership, Invested and Realized are post-activity. "
                                "Operating metrics are carried unchanged from the source file and should be refreshed."])
    ws.append([f"{next_label} Activity tab", f"Empty. Record every portfolio event in {next_label} here, then run hc-valuation "
                                             "with the matching policy file."])
    ws.append(["Open Items tab", "Unresolved items carried from the prior run (also emitted as open_items_carry.yaml, "
                                 "which the pipeline reads as prior_open_items)."])
    ws.append(["Snapshot Notes tab", "Rows whose Prior Mark deliberately departs from Ownership x Latest Post-Money "
                                     "(note at cost, pending deal, override). Also carried in open_items_carry.yaml as mark_basis."])
    for c in ws[1]:
        c.font = Font(bold=True)


SYNTHETIC_SHEET = "SYNTHETIC TEST DATA"   # the marker a synthetic workbook carries as its first sheet (workbooks.py)


def _carry_synthetic_marker(wb: Workbook, source_path: Path) -> None:
    """Invented test data stays labelled through a roll-forward: if the source workbook carries the
    marker sheet, the emitted book carries it too, first. Without this the next quarter's input
    would look like a real book and be offered as one (found when the synthetic chain's emitted
    inputs appeared in the real book's Workbook select)."""
    src_wb = openpyxl.load_workbook(source_path, read_only=True)
    name = next((n for n in src_wb.sheetnames if n.strip().upper() == SYNTHETIC_SHEET), None)
    rows = [list(r) for r in src_wb[name].iter_rows(values_only=True)] if name else []
    src_wb.close()
    if name is None:
        return
    ws = wb.create_sheet(SYNTHETIC_SHEET, 0)
    for row in rows or [[SYNTHETIC_SHEET]]:
        ws.append(row)
    ws.append(["This book was emitted by hc-valuation build from a synthetic quarter; it is invented test data too."])
    ws.column_dimensions["A"].width = 120


def _write_open_items(wb: Workbook, run: ValuationRun) -> None:
    ws = wb.create_sheet("Open Items")
    ws.append(["Company", "Kind", "Opened", "Opened Quarter", "Expected Resolution", "Amount ($M)", "Detail", "Age (quarters)", "Escalated"])
    for o in run.open_items:
        ws.append([o.company, o.kind.value, o.opened, o.opened_quarter, o.expected_resolution, o.amount_musd, o.detail,
                   o.age_quarters, "yes" if o.escalated else "no"])
    for col, w in zip("ABCDEFGHI", (22, 20, 12, 14, 18, 12, 70, 14, 10)):
        ws.column_dimensions[col].width = w
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
    _style_header(ws, 9, "A2")


def _write_notes(wb: Workbook, notes: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet("Snapshot Notes")
    ws.append(["Company", "Note"])
    for company, note in notes:
        ws.append([company, note])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120
    _style_header(ws, 2, "A2")


def note_legs(run: ValuationRun) -> list[dict[str, Any]]:
    """Positions whose booked mark carries a convertible-note leg at cost (M-060). The Portfolio
    tab has one `Prior Mark` column, so next quarter the leg would otherwise be read as equity —
    and a note repaid in that quarter would then be counted twice: the cash as realized and the
    principal still inside the mark (found by the synthetic Q1 2027 chain, HARDENING_REPORT.md
    D-5). The sidecar carries the leg as a number the next run seeds `note_at_cost` from."""
    out = []
    for c in run.companies:
        if c.status_after == Status.ACTIVE and c.note_at_cost > 0:
            out.append({"company": c.company, "amount_musd": round(c.note_at_cost, 6),
                        "reason": f"note leg ${c.note_at_cost:.2f}M carried at cost inside the ${c.booked_mark:.2f}M prior mark (M-060)"})
    return out


def write_open_items_sidecar(run: ValuationRun, path: str | Path, mark_basis: list[tuple[str, str]] | None = None,
                             staleness_anchors: list[dict[str, str]] | None = None) -> Path:
    """`open_items_carry.yaml`: what the pipeline loads next quarter — `open_items` as
    `prior_open_items`, `mark_basis` as the explained departures X-904 accepts,
    `staleness_anchors` for clocks that outlive the `Latest Round` column, and `note_legs` for
    the part of a prior mark that is a note at cost rather than equity."""
    p = Path(path)
    payload = {
        "source_run_id": run.manifest.run_id,
        "quarter": run.manifest.quarter_label,
        "open_items": [o.model_dump(mode="json") for o in run.open_items],
        "mark_basis": [{"company": c, "reason": r} for c, r in (mark_basis or [])],
        "staleness_anchors": list(staleness_anchors or []),
        "note_legs": note_legs(run),
    }
    p.write_text(yaml.safe_dump(payload, sort_keys=False, allow_unicode=True))
    return p


def write_next_quarter_workbook(run: ValuationRun, source_workbook_path: str | Path, out_path: str | Path,
                                cfg: RuleConfig) -> Path:
    """Emit the next quarter's input workbook plus the `open_items_carry.yaml` sidecar beside it."""
    source = Path(source_workbook_path)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    next_label = next_quarter_label(cfg.quarter.label)
    activity_name = f"{next_label} Activity"
    if not re.match(cfg.schema_.activity_sheet_pattern, activity_name):
        raise ValueError(f"emitted activity sheet name {activity_name!r} would not match the policy pattern "
                         f"{cfg.schema_.activity_sheet_pattern!r}; the file could not be re-ingested")

    snapshot, _feed = read_workbook(source, cfg)
    sources = snapshot.by_company()
    tol = cfg.tolerances.prior_mark_reconciliation_musd

    wb = Workbook()
    notes = _write_portfolio(wb, run, sources, tol)
    _write_activity(wb, activity_name)
    _copy_field_definitions(wb, source, run, next_label)
    _write_open_items(wb, run)
    anchors = _carried_anchors(run, sources)
    _write_notes(wb, notes + [(a["company"], a["reason"]) for a in anchors])
    _carry_synthetic_marker(wb, source)
    wb.save(out)
    write_open_items_sidecar(run, out.parent / "open_items_carry.yaml", mark_basis=notes, staleness_anchors=anchors)
    return out
