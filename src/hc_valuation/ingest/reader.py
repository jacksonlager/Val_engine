"""Workbook -> (PortfolioSnapshot, ActivityFeed).

The activity sheet is located by regex, never by a literal name: this engine is meant
to run in quarters that do not exist yet. Everything the reader tolerates on the way —
a header on row 3, `Post Money` for `Post-Money`, `$28.2M` in a number cell, `Nimbrel
Inc.` for `Nimbrel` — is recorded as a `Correction` on the snapshot / feed so that
`validate.py` can show the reviewer what was read as what. The reader never guesses:
an ambiguous cell is passed through with a blocking correction, and a file it cannot
read at all raises `IngestError` with a message that names the fix.
"""
from __future__ import annotations

import hashlib
import re
import zipfile
from datetime import date
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from ..config import NormalizationCfg, RuleConfig
from ..engine.inputs import Correction, EventType
from . import normalize as nz
from .schema import (
    ACTIVITY_COLUMNS, ACTIVITY_HEADER_ALIASES, GROWTH_COLUMNS, MUSD_COLUMNS, PERCENT_COLUMNS,
    PORTFOLIO_COLUMNS, PORTFOLIO_HEADER_ALIASES, REQUIRED_ACTIVITY, REQUIRED_PORTFOLIO,
    ActivityFeed, Event, PortfolioSnapshot, Position, Status,
)


class IngestError(ValueError):
    """A named, actionable failure in the input file."""


Row = Sequence[Any]


def file_sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    h.update(Path(path).read_bytes())
    return h.hexdigest()


# ---------------------------------------------------------------- opening the file (SPEC §2.5)

def _describe_bytes(p: Path) -> str:
    """What the file actually is, for an error that says what was found."""
    head = p.read_bytes()[:512]
    if not head:
        return "an empty file"
    if head.startswith(b"PK"):
        return "a zip container without workbook parts"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "a legacy .xls (OLE) file"
    try:
        text = head.decode("utf-8")
    except UnicodeDecodeError:
        return "binary data that is not a zip container"
    return "text that looks like CSV" if "," in text else "plain text"


def _open_workbook(p: Path):
    """openpyxl's failures become one IngestError that says what was found and what was
    expected — a reviewer must never see a zipfile traceback."""
    if not p.exists():
        raise IngestError(f"input workbook not found: {p}")
    try:
        return openpyxl.load_workbook(p, data_only=True, read_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as ex:
        raise IngestError(
            f"{p.name} is not a workbook: found {_describe_bytes(p)}; expected an .xlsx file "
            f"(an Excel 2007+ zip container with a Portfolio tab and a quarterly Activity tab). "
            f"Underlying error: {type(ex).__name__}: {ex}"
        ) from None


# ---------------------------------------------------------------- sheet discovery (SPEC §2.1)

def find_activity_sheet(sheetnames: list[str], pattern: str) -> tuple[str, Correction | None]:
    """The policy regex decides first. When nothing matches it, any sheet whose folded name
    contains `activity` or `events` is a candidate: exactly one is used with X-919; several
    are read as the first with a blocking X-914 that names them all, so the run cannot
    produce a number from a guess."""
    rx = re.compile(pattern)
    matches = [s for s in sheetnames if rx.match(s)]
    if len(matches) > 1:
        raise IngestError(f"more than one activity sheet matches /{pattern}/: {matches}. Keep one per file.")
    if matches:
        return matches[0], None
    relaxed = nz.relaxed_activity_candidates(sheetnames)
    if not relaxed:
        raise IngestError(
            f"no activity sheet matches /{pattern}/ and no sheet name contains 'activity' or 'events' — "
            f"sheets present: {sheetnames}. Expected a tab named like 'Q4 2026 Activity'."
        )
    if len(relaxed) > 1:
        return relaxed[0], Correction(kind="ambiguous", original=pattern, resolved=relaxed[0], method="sheet",
                                      detail="more than one activity-like sheet: " + "; ".join(relaxed),
                                      sheet=relaxed[0])
    return relaxed[0], Correction(kind="sheet", original=pattern, resolved=relaxed[0], method="relaxed",
                                  sheet=relaxed[0])


def _find_portfolio_sheet(sheetnames: list[str], exact: str) -> tuple[str, Correction | None]:
    name, corr, hits = nz.find_sheet(sheetnames, exact)
    if name is not None:
        return name, (corr.located(sheet=name) if corr is not None else None)
    if len(hits) > 1:
        raise IngestError(f"portfolio sheet '{exact}' not found and several sheets look like a book: {hits}. "
                          f"Name one of them '{exact}'.")
    raise IngestError(f"portfolio sheet '{exact}' not found; sheets: {sheetnames}. "
                      f"Expected a tab named '{exact}' (or portfolio / book / positions / holdings).")


def read_workbook(path: str | Path, config: RuleConfig) -> tuple[PortfolioSnapshot, ActivityFeed]:
    p = Path(path)
    wb = _open_workbook(p)
    sc = config.schema_
    try:
        port_name, port_corr = _find_portfolio_sheet(wb.sheetnames, sc.portfolio_sheet_name)
        act_name, act_corr = find_activity_sheet(wb.sheetnames, sc.activity_sheet_pattern)
        snapshot = _read_portfolio(wb[port_name], config, [port_corr] if port_corr else [])
        feed = _read_activity(wb[act_name], act_name, config, snapshot, [act_corr] if act_corr else [])
    finally:
        wb.close()
    return snapshot, feed


# ---------------------------------------------------------------- headers and rows (SPEC §2.2, §2.4)

class _Sheet:
    """A sheet after header discovery: the column index, the raw headers, the data rows
    (each tagged with its 1-based Excel row), and the corrections made getting there."""

    def __init__(self, ws, columns: Mapping[str, str], aliases: Mapping[str, str], required: Sequence[str],
                 config: RuleConfig, corrections: list[Correction]) -> None:
        self.title: str = ws.title
        self.corrections = corrections
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        ncfg = config.normalization
        found = nz.find_header_row(rows, list(columns), aliases, ncfg)
        if found is None:
            raise IngestError(
                f"sheet '{self.title}' has no header row in its first {ncfg.header_scan_rows} rows: no row carries "
                f"{ncfg.header_min_known_columns} or more known column names. Expected columns like {list(columns)[:4]}."
            )
        self.header_row, self.idx, raw_headers, header_corrs = found
        self.corrections.extend(c.located(sheet=self.title) for c in header_corrs)
        if self.header_row != 1:
            self.corrections.append(Correction(
                kind="structure", original=f"row {self.header_row}", resolved="header", method="header-row",
                detail=f"header found on row {self.header_row}; {self.header_row - 1} row(s) above it ignored",
                sheet=self.title, row_index=self.header_row))
        self._check_required(required, raw_headers)
        resolved = {raw for raw in raw_headers if self._resolves(raw, columns, aliases, ncfg)}
        self.unknown: tuple[str, ...] = tuple(h for h in raw_headers if h not in resolved)
        if self.unknown and config.schema_.unknown_column == "fail":
            raise IngestError(f"unknown column(s) in {self.title}: {self.unknown}")
        header_row = rows[self.header_row - 1]
        self.unknown_idx: dict[str, int] = {
            nz.unify_typography(str(v)): i for i, v in enumerate(header_row)
            if v is not None and nz.unify_typography(str(v)) in self.unknown}
        self.rows: list[tuple[int, Row]] = self._data_rows(rows)

    @staticmethod
    def _resolves(raw: str, columns: Mapping[str, str], aliases: Mapping[str, str], ncfg: NormalizationCfg) -> bool:
        canon, _ = nz.normalize_header(raw, list(columns), aliases, ncfg)
        return canon is not None

    def _check_required(self, required: Sequence[str], raw_headers: Sequence[str]) -> None:
        missing = [c for c in required if c not in self.idx]
        if not missing:
            return
        hints = []
        for col in missing:
            near = nz.closest_header(col, raw_headers)
            hints.append(f"{col!r}" + (f" (closest header seen: {near!r})" if near else ""))
        raise IngestError(f"sheet '{self.title}' is missing required column(s): {missing}. " + "; ".join(hints))

    def _data_rows(self, rows: Sequence[Row]) -> list[tuple[int, Row]]:
        """Rows below the header, minus blanks, totals and trailing notes — each skip recorded
        (blank rows once per sheet, with the count). Trailing blank rows are Excel's dimension
        slop, not structure, and are dropped silently."""
        body = list(enumerate(rows[self.header_row:], start=self.header_row + 1))
        while body and nz.row_is_blank(body[-1][1]):
            body.pop()
        kept: list[tuple[int, Row]] = []
        blanks = 0
        for r, row in body:
            if nz.row_is_blank(row):
                blanks += 1
            elif nz.row_is_total(row):
                self.corrections.append(Correction(kind="structure", original=str(row[0] or next(v for v in row if v)),
                                                   resolved="skipped", method="total-row",
                                                   detail=f"row {r} is a totals row; skipped", sheet=self.title, row_index=r))
            elif nz.row_is_note(row):
                self.corrections.append(Correction(kind="structure", original=str(row[0]), resolved="skipped",
                                                   method="note-row", detail=f"row {r} carries only a note; skipped",
                                                   sheet=self.title, row_index=r))
            else:
                kept.append((r, row))
        if blanks:
            self.corrections.append(Correction(kind="structure", original=f"{blanks} blank row(s)", resolved="skipped",
                                               method="blank-rows", detail=f"{blanks} blank row{'s' if blanks != 1 else ''} inside the data skipped",
                                               sheet=self.title))
        return kept

    def getter(self, row: Row) -> Callable[[str], Any]:
        return lambda col: row[self.idx[col]] if col in self.idx and self.idx[col] < len(row) else None

    def extras(self, row: Row) -> dict[str, Any]:
        """Unknown columns are carried verbatim, keyed by their raw header."""
        return {h: row[i] for h, i in self.unknown_idx.items() if i < len(row)}


class _Cells:
    """Per-row coercion that records every correction against the right cell."""

    def __init__(self, sheet: _Sheet, row_index: int, get: Callable[[str], Any], ncfg: NormalizationCfg,
                 company: str | None = None) -> None:
        self.sheet, self.row_index, self.get, self.ncfg, self.company = sheet, row_index, get, ncfg, company

    def _record(self, corrs: Sequence[Correction], column: str) -> None:
        self.sheet.corrections.extend(
            c.located(sheet=self.sheet.title, row_index=self.row_index, column=column, company=self.company) for c in corrs)

    def number(self, column: str) -> float | None:
        raw = self.get(column)
        if column in MUSD_COLUMNS:
            value, corrs = nz.coerce_musd(raw, self.ncfg)
        elif column in PERCENT_COLUMNS:
            value, corrs = nz.coerce_percent(raw, self.ncfg)
        elif column in GROWTH_COLUMNS:
            value, corrs = nz.coerce_percent(raw, self.ncfg, points_heuristic=False)
        else:
            value, corr = nz.coerce_number(raw)
            corrs = [corr] if corr is not None else []
        self._record(corrs, column)
        return value

    def day(self, column: str, required: bool = True) -> date | None:
        raw = self.get(column)
        try:
            value, corr = nz.coerce_date(raw)
        except nz.DateParseError as ex:
            raise IngestError(f"sheet '{self.sheet.title}' row {self.row_index}: {column!r} is not a date: {ex}") from None
        if corr is not None:
            self._record([corr], column)
        if value is None and required:
            raise IngestError(f"sheet '{self.sheet.title}' row {self.row_index}: {column!r} is blank; a date is required")
        return value

    def text(self, column: str) -> str:
        raw = self.get(column)
        return "" if raw is None else nz.unify_typography(str(raw))


# ---------------------------------------------------------------- Portfolio tab

_STATUS_SYNONYMS: dict[str, Status] = {
    "exited": Status.ACQUIRED, "realized": Status.ACQUIRED, "realised": Status.ACQUIRED, "sold": Status.ACQUIRED,
    "acquired (stock)": Status.ACQUIRED, "merged": Status.ACQUIRED,
    "public": Status.ACTIVE, "listed": Status.ACTIVE, "ipo": Status.ACTIVE, "live": Status.ACTIVE, "open": Status.ACTIVE, "held": Status.ACTIVE,
    "shutdown": Status.SHUT_DOWN, "closed": Status.SHUT_DOWN, "dissolved": Status.SHUT_DOWN, "liquidated": Status.SHUT_DOWN,
    "wound down": Status.SHUT_DOWN, "written off": Status.SHUT_DOWN, "write-off": Status.SHUT_DOWN, "writeoff": Status.SHUT_DOWN,
    "bankrupt": Status.SHUT_DOWN, "ceased operations": Status.SHUT_DOWN, "defunct": Status.SHUT_DOWN,
}


def _status(raw: Any, cells: _Cells, company: str) -> Status:
    """Exact status, else a case/whitespace-folded one, else a known synonym (both recorded). A word
    the engine does not know is that row's problem, not the workbook's: the position is read as
    Active and a blocking correction refuses its row (X-925) — the other ninety-nine still value."""
    text = "" if raw is None else str(raw)
    try:
        return Status(text)
    except ValueError:
        pass
    for member in Status:
        if nz.fold(text) == nz.fold(member.value):
            cells._record([Correction(kind="value", original=text, resolved=member.value, method="fold")], "Status")
            return member
    syn = _STATUS_SYNONYMS.get(nz.fold(text))
    if syn is not None:
        prior = cells.number("Prior Mark ($M)") or 0.0
        if syn is Status.ACTIVE or prior <= 0.0:
            cells._record([Correction(kind="value", original=text, resolved=syn.value, method="synonym")], "Status")
            return syn
        cells._record([Correction(kind="status", original=text.strip(), resolved=Status.ACTIVE.value, method="synonym-with-mark",
                                  detail=f"Status {text.strip()!r} reads as {syn.value}, but the row still carries a ${prior:.2f}M prior mark; "
                                         "held as Active and blocked until the status or the mark is corrected")], "Status")
        return Status.ACTIVE
    cells._record([Correction(kind="status", original=text.strip(), resolved=Status.ACTIVE.value, method="unknown",
                              detail=f"Status {text.strip()!r} is not Active, Acquired or Shut Down; the row is held as Active and blocked until it is corrected")], "Status")
    return Status.ACTIVE


def _read_portfolio(ws, config: RuleConfig, corrections: list[Correction]) -> PortfolioSnapshot:
    sheet = _Sheet(ws, PORTFOLIO_COLUMNS, PORTFOLIO_HEADER_ALIASES, REQUIRED_PORTFOLIO, config, corrections)
    ncfg = config.normalization

    positions: list[Position] = []
    for r, row in sheet.rows:
        get = sheet.getter(row)
        company, corr = nz.clean_company(get("Company"))
        if not company:
            continue
        cells = _Cells(sheet, r, get, ncfg, company)
        if corr is not None:
            cells._record([corr], "Company")
        status = _status(get("Status"), cells, company)
        first_inv = cells.day("First Investment", required=False)
        latest_round = cells.day("Latest Round", required=False)
        if latest_round is None:
            # A blank round date is that row's problem: carry the first-investment date (recorded) or, with
            # neither, the prior close, and refuse the row (X-925) so a person fills the cell in.
            if first_inv is not None:
                cells._record([Correction(kind="value", original="", resolved=first_inv.isoformat(), method="first-investment",
                                          detail="Latest Round blank; the First Investment date stands in")], "Latest Round")
                latest_round = first_inv
            else:
                latest_round = config.quarter.prior_close
                cells._record([Correction(kind="status", original="", resolved=latest_round.isoformat(), method="blank",
                                          detail="Latest Round and First Investment are both blank; the prior close stands in and the row is blocked until a date is entered")], "Latest Round")
        positions.append(Position(
            company=company,
            sector=cells.text("Sector"),
            fund=cells.text("Fund"),
            stage=cells.text("Stage"),
            status=status,
            first_investment=first_inv if first_inv is not None else latest_round,
            latest_round=latest_round,
            latest_post_money=cells.number("Latest Post-Money ($M)") or 0.0,
            invested=cells.number("Invested ($M)") or 0.0,
            ownership=cells.number("Ownership (FD %)") or 0.0,
            prior_mark=cells.number("Prior Mark ($M)") or 0.0,
            realized=cells.number("Realized ($M)") or 0.0,
            arr=cells.number("ARR ($M)"),
            arr_growth=cells.number("ARR Growth (YoY %)"),
            gross_margin=cells.number("Gross Margin (%)"),
            net_burn=cells.number("Net Burn ($M/mo)"),
            cash=cells.number("Cash ($M)"),
            headcount=int(cells.number("Headcount") or 0) or None,
            sheet_moic=cells.number("MOIC (x)"),
            sheet_runway=cells.number("Runway (mo)"),
            row_index=r,
            extra=sheet.extras(row),
        ))
    if not positions:
        raise IngestError(f"sheet '{sheet.title}' has a header but no portfolio rows; the book cannot be empty.")
    return PortfolioSnapshot(
        as_of=config.quarter.prior_close, positions=tuple(positions),
        sheet_name=sheet.title, unknown_columns=sheet.unknown, corrections=tuple(sheet.corrections),
    )


# ---------------------------------------------------------------- Activity tab

def _read_activity(ws, sheet_name: str, config: RuleConfig, snapshot: PortfolioSnapshot,
                   corrections: list[Correction]) -> ActivityFeed:
    sheet = _Sheet(ws, ACTIVITY_COLUMNS, ACTIVITY_HEADER_ALIASES, REQUIRED_ACTIVITY, config, corrections)
    ncfg = config.normalization
    book = [p.company for p in snapshot.positions]

    events: list[Event] = []
    for r, row in sheet.rows:
        get = sheet.getter(row)
        raw_company = get("Company")
        if raw_company is None or str(raw_company).strip() == "":
            if any(get(col) not in (None, "") for col in ("Event", "Post-Money / Deal Value ($M)", "HC Investment ($M)", "Proceeds to HC ($M)")):
                corrections.append(Correction(kind="ambiguous", original="", resolved="", method="company", sheet=ws.title, row_index=r,
                                              column="Company", detail=f"row {r} carries an event but no company name; it cannot be attached to a position"))
            continue
        company, ccorr = nz.match_company(raw_company, book, ncfg)
        cells = _Cells(sheet, r, get, ncfg, company)
        if ccorr is not None:
            cells._record([ccorr], "Company")
        raw_type = get("Event")
        event_type, ecorr = nz.normalize_event_type(raw_type, ncfg)
        if ecorr is not None:
            cells._record([ecorr], "Event")
        extra = sheet.extras(row)
        # A priced round on a company the Portfolio tab does not have, with HC's cheque on the row,
        # is HC's first investment in it — the row `New Investment` describes, filed under the round's
        # name. Read it as one (recorded as X-912, like any other event-type reading) so the position
        # is created (M-014, X-918) instead of the row being refused as an unknown company (X-901).
        # Found by the Q2 2026 test file, whose two new holdings arrived exactly this way.
        if (company not in book and event_type == EventType.PRICED_ROUND.value
                and (cells.number("HC Investment ($M)") or 0) > 0):
            cells._record([Correction(kind="event_type", original=str(raw_type), resolved=EventType.NEW_INVESTMENT.value, method="context",
                                      detail=f"{company} is not in the Portfolio tab and HC's cheque is on the row: read as an initial investment")], "Event")
            event_type = EventType.NEW_INVESTMENT.value
        if raw_type is not None and str(raw_type) != event_type:
            extra["raw_event_type"] = str(raw_type)
        if str(raw_company) != company:
            extra["raw_company"] = str(raw_company)
        # A blank or unreadable date refuses this row (X-902, below in validate) rather than the
        # whole workbook: the measurement date stands in only so the row can be carried through
        # to the refusal, and `extra` says why, so no window check reads the stand-in as a fact.
        try:
            when = cells.day("Date", required=False)
        except IngestError as ex:
            when = None
            extra["date_unreadable"] = str(ex).split(": ", 2)[-1]
        if when is None:
            extra.setdefault("date_missing", True)
            when = config.quarter.measurement_date
        events.append(Event(
            date=when,
            company=company,
            event_type=event_type,
            detail=cells.text("Detail"),
            value=cells.number("Post-Money / Deal Value ($M)"),
            hc_investment=cells.number("HC Investment ($M)"),
            ownership_after=cells.number("HC Ownership After (FD %)"),
            proceeds=cells.number("Proceeds to HC ($M)"),
            notes=cells.text("Notes"),
            row_index=r,
            extra=extra,
        ))
    label = nz.parse_quarter_label(sheet_name) or config.quarter.label
    return ActivityFeed(quarter_label=label, events=tuple(events), sheet_name=sheet_name,
                        unknown_columns=sheet.unknown, corrections=tuple(sheet.corrections))


__all__ = ["IngestError", "read_workbook", "find_activity_sheet", "file_sha256"]
