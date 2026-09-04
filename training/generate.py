"""Deterministic scenario workbook generator: scenarios/*.yaml -> workbooks/*.xlsx.

    python training/generate.py            # every scenario
    python training/generate.py 06_value_formats 08_structure

Every workbook is rebuilt from scratch (never edited in place) so the dirty representation a
scenario asks for is exactly what lands in the file: a string like "$28.2M" stays a string,
a YAML date becomes an Excel datetime, an integer in the Date column stays an Excel serial.
Nothing here is random except the seeded synthesiser behind `19_big_book`.

Mutation order (SPEC §4): base rows -> portfolio_mutations -> activity rows replaced ->
header_mutations / drop_columns / extra_columns / column_order -> structure (title rows,
blank rows, totals row, trailing note) -> sheet_names / extra_sheets.
"""
from __future__ import annotations

import csv
import io
import random
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font

TRAINING = Path(__file__).resolve().parent
ROOT = TRAINING.parent
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from hc_valuation.ingest.schema import ACTIVITY_COLUMNS, PORTFOLIO_COLUMNS  # noqa: E402

SCENARIOS = TRAINING / "scenarios"
WORKBOOKS = TRAINING / "workbooks"
BASE_WORKBOOK = ROOT / "data" / "HC_Mock_Portfolio_Data.xlsx"
BASE_ACTIVITY_SHEET = "Q3 2026 Activity"

# scenario activity keys -> workbook headers (attr names from ingest/schema.py, plus `event`)
_EVT_KEY_TO_HEADER: dict[str, str] = {attr: hdr for hdr, attr in ACTIVITY_COLUMNS.items()}
_EVT_KEY_TO_HEADER["event"] = "Event"

SECTORS = ["AI/ML", "Fintech", "Developer Tools", "Climate & Energy", "Cybersecurity", "Infrastructure",
           "Enterprise SaaS", "Healthcare", "Space & Defense", "Data & Analytics", "Robotics", "Consumer"]
FUNDS = ["Fund I", "Fund II", "Fund III"]
STAGES = ["Seed", "Series A", "Series B", "Series C", "Series D+"]


# ----------------------------------------------------------------------------- loading

def load_scenario(path: Path) -> dict[str, Any]:
    raw = yaml.safe_load(path.read_text())
    raw.setdefault("name", path.stem)
    return raw


def _sheet_rows(ws) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


class Book:
    """A workbook held as plain rows: {sheet name: [row, ...]} plus the two tabs we mutate as
    header + list-of-dicts."""

    def __init__(self, portfolio_header: list[str], portfolio: list[dict[str, Any]],
                 activity_sheet: str, others: dict[str, list[list[Any]]]) -> None:
        self.portfolio_header = portfolio_header
        self.portfolio = portfolio
        self.activity_sheet = activity_sheet
        self.activity_header = list(ACTIVITY_COLUMNS)
        self.activity: list[dict[str, Any]] = []
        self.others = others            # verbatim sheets (Field Definitions, Open Items, ...)
        self.sidecar: Path | None = None


def _load_base(path: Path, activity_sheet: str) -> Book:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Portfolio"]
    rows = _sheet_rows(ws)
    header = [str(h) for h in rows[0]]
    portfolio = [dict(zip(header, r)) for r in rows[1:] if r and r[0] is not None]
    others = {name: _sheet_rows(wb[name]) for name in wb.sheetnames if name not in ("Portfolio", activity_sheet)}
    wb.close()
    return Book(header, portfolio, activity_sheet, others)


def _snapshot_base(tmp: Path) -> tuple[Path, Path]:
    """Run Q3 on the real workbook and emit the next-quarter book + sidecar into `tmp`."""
    from hc_valuation import pipeline
    from hc_valuation.export.snapshot import write_next_quarter_workbook

    paths = pipeline.RunPaths.default(root=ROOT, workbook=BASE_WORKBOOK, policy=ROOT / "rules" / "2026Q3.yaml")
    res = pipeline.execute(paths, adjudicate=False, generated_at=datetime(2026, 9, 30), provider="stub")
    out = write_next_quarter_workbook(res.run, BASE_WORKBOOK, tmp / "snapshot_2026Q4.xlsx", res.config)
    return out, out.parent / "open_items_carry.yaml"


# ----------------------------------------------------------------------------- mutations

def _apply_portfolio_mutations(book: Book, muts: dict[str, dict[str, Any]] | None) -> None:
    if not muts:
        return
    by_name = {p["Company"]: p for p in book.portfolio}
    for company, changes in muts.items():
        if changes is None:           # `Company: null` drops the row
            book.portfolio = [p for p in book.portfolio if p["Company"] != company]
            continue
        if company in by_name:
            by_name[company].update(changes)
        else:                          # a company the book does not have: append a full row
            row = {h: None for h in book.portfolio_header}
            row["Company"] = company
            row.update(changes)
            book.portfolio.append(row)


def _activity_rows(rows: list[dict[str, Any]] | None) -> tuple[list[str], list[dict[str, Any]]]:
    header = list(ACTIVITY_COLUMNS)
    out: list[dict[str, Any]] = []
    for r in rows or []:
        row: dict[str, Any] = {h: None for h in header}
        for k, v in r.items():
            hdr = _EVT_KEY_TO_HEADER.get(k, k)
            if hdr not in row:
                header.append(hdr)          # an extra column named in a row
                for prev in out:
                    prev.setdefault(hdr, None)
                row[hdr] = None
            row[hdr] = v
        out.append(row)
    return header, out


def _mutate_headers(header: list[str], rows: list[dict[str, Any]], spec: dict[str, Any], tab: str
                    ) -> tuple[list[str], list[dict[str, Any]]]:
    """Rename / drop / add / reorder columns. Rows keep their original keys; a `label` map
    carries the header text that is actually written."""
    drops = set((spec.get("drop_columns") or {}).get(tab, []) or [])
    header = [h for h in header if h not in drops]
    extra = (spec.get("extra_columns") or {}).get(tab, {}) or {}
    for h, v in extra.items():
        if h not in header:
            header.append(h)
        for r in rows:
            r.setdefault(h, v)
    order = spec.get("column_order") if tab == "activity" else (spec.get("portfolio_column_order"))
    if order:
        missing = [h for h in order if h not in header]
        if missing:
            raise ValueError(f"column_order names columns not in the {tab} tab: {missing}")
        header = list(order) + [h for h in header if h not in order]
    return header, rows


def _labels(header: list[str], spec: dict[str, Any], tab: str) -> list[Any]:
    renames = (spec.get("header_mutations") or {}).get(tab, {}) or {}
    return [renames.get(h, h) for h in header]


def _write_tab(wb: Workbook, name: str, header: list[str], labels: list[Any], rows: list[dict[str, Any]],
               structure: dict[str, Any] | None, first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    structure = structure or {}
    ncols = max(len(header), 1)

    title_rows = int(structure.get("title_rows", 0) or 0)
    titles = structure.get("titles") or ["HC Portfolio Workbook", "Prepared by HC Finance — do not edit"]
    for i in range(title_rows):
        ws.append([titles[i] if i < len(titles) else None] + [None] * (ncols - 1))
        if i == 0 and ncols > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
            ws.cell(row=1, column=1).font = Font(bold=True, size=13)
    ws.append(labels)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    every = int(structure.get("blank_rows_every", 0) or 0)
    sums: dict[str, float] = {}
    for i, r in enumerate(rows, start=1):
        ws.append([_cell(r.get(h)) for h in header])
        for h in header:
            v = r.get(h)
            if isinstance(v, (int, float)) and not isinstance(v, bool) and not isinstance(v, datetime):
                sums[h] = sums.get(h, 0.0) + float(v)
        if every and i % every == 0:
            ws.append([None] * ncols)
    if structure.get("totals_row"):
        total = [None] * ncols
        total[0] = structure.get("totals_label", "Total")
        for j, h in enumerate(header[1:], start=1):
            if h in sums and ("$M" in h or h in ("Headcount",)):
                total[j] = round(sums[h], 4)
        ws.append(total)
    note = structure.get("trailing_note")
    if note:
        ws.append([note] + [None] * (ncols - 1))


def _cell(v: Any) -> Any:
    """Write values exactly as the scenario gives them. A YAML date becomes an Excel datetime
    (that is what a real cell holds); strings, ints, floats and None pass through untouched."""
    if isinstance(v, date) and not isinstance(v, datetime):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, dict) and "repeat" in v:      # {repeat: "text ", times: 400} -> a very long cell
        return str(v["repeat"]) * int(v.get("times", 1))
    return v


def _save(wb: Workbook, out: Path) -> None:
    """Pin the document properties and the zip entry timestamps so regenerating a workbook is
    byte-identical (openpyxl stamps both with the wall clock)."""
    fixed = datetime(2026, 9, 30)
    wb.properties.creator = "hc-valuation training corpus"
    wb.properties.lastModifiedBy = "generate.py"
    wb.properties.created = fixed
    wb.properties.modified = fixed
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed_info = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 30, 0, 0, 0))
            fixed_info.compress_type = zipfile.ZIP_DEFLATED
            fixed_info.external_attr = info.external_attr
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":   # openpyxl re-stamps `modified` at save time
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-09-30T00:00:00Z\g<2>", data)
            dst.writestr(fixed_info, data)


def _write_others(wb: Workbook, others: dict[str, list[list[Any]]]) -> None:
    for name, rows in others.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append([_cell(v) for v in r])


# ----------------------------------------------------------------------------- builders

def build_standard(spec: dict[str, Any], out: Path) -> Path:
    base = str(spec.get("base", BASE_WORKBOOK.name))
    tmp: Path | None = None
    if base.startswith("snapshot:"):
        tmp = Path(tempfile.mkdtemp(prefix="hc_snapshot_"))
        src, sidecar = _snapshot_base(tmp)
        book = _load_base(src, "Q4 2026 Activity")
        book.activity_sheet = "Q4 2026 Activity"
        shutil.copy(sidecar, out.with_name(out.stem + ".open_items_carry.yaml"))
    else:
        src = BASE_WORKBOOK if base == BASE_WORKBOOK.name else (ROOT / "data" / base)
        book = _load_base(src, BASE_ACTIVITY_SHEET)

    _apply_portfolio_mutations(book, spec.get("portfolio_mutations"))
    for name in spec.get("portfolio_duplicates") or []:      # a company row pasted twice into the book
        src_row = next(p for p in book.portfolio if p["Company"] == name)
        book.portfolio.append(dict(src_row))
    act_header, act_rows = _activity_rows(spec.get("activity"))
    act_header, act_rows = _mutate_headers(act_header, act_rows, spec, "activity")
    pf_header, pf_rows = _mutate_headers(list(book.portfolio_header), book.portfolio, spec, "portfolio")

    names = spec.get("sheet_names") or {}
    pf_name = names.get("portfolio", "Portfolio")
    act_name = names.get("activity", book.activity_sheet)
    structure = spec.get("structure") or {}
    # `structure` may be flat (applies to the activity tab) or keyed by tab
    if structure and not ({"activity", "portfolio"} & set(structure)):
        structure = {"activity": structure}

    wb = Workbook()
    _write_tab(wb, pf_name, pf_header, _labels(pf_header, spec, "portfolio"), pf_rows, structure.get("portfolio"), first=True)
    _write_tab(wb, act_name, act_header, _labels(act_header, spec, "activity"), act_rows, structure.get("activity"))
    for extra in spec.get("extra_sheets") or []:
        if extra.get("copy_of", "activity") == "activity":
            _write_tab(wb, extra["name"], act_header, _labels(act_header, spec, "activity"), act_rows, structure.get("activity"))
        else:
            _write_tab(wb, extra["name"], pf_header, _labels(pf_header, spec, "portfolio"), pf_rows, structure.get("portfolio"))
    _write_others(wb, book.others)
    _save(wb, out)
    if tmp is not None:
        shutil.rmtree(tmp, ignore_errors=True)
    return out


def synthesise(spec: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """`19_big_book`: N coherent positions (mark = ownership × post, so X-904 passes) and K
    events cycling the eight original types on distinct companies. Seeded; no clock."""
    syn = spec.get("synthesize") or {}
    n = int(syn.get("companies", 1000))
    k = int(syn.get("events", 120))
    rng = random.Random(int(syn.get("seed", 19)))
    positions: list[dict[str, Any]] = []
    for i in range(1, n + 1):
        post = round(8.0 + (i * 7.3) % 400.0, 1)   # floor keeps ARR (= post/12) above the X-403 screen
        own = round(0.02 + (i * 0.0137) % 0.12, 4)
        mark = round(own * post, 6)
        invested = round(mark / 1.8, 6)
        arr = round(post / 12.0, 3)
        burn = round(arr / 20.0 + 0.05, 3)
        latest = date(2025, 1, 1) + timedelta(days=(i * 37) % 500)
        status = "Active"
        if i % 97 == 0:
            status = "Acquired" if i % 2 else "Shut Down"
        row = {
            "Company": f"Synth-{i:04d}", "Sector": SECTORS[i % len(SECTORS)], "Fund": FUNDS[i % len(FUNDS)],
            "Stage": STAGES[i % len(STAGES)], "Status": status,
            "First Investment": latest - timedelta(days=300), "Latest Round": latest,
            "Latest Post-Money ($M)": post, "Invested ($M)": invested, "Ownership (FD %)": own,
            "Prior Mark ($M)": mark if status == "Active" else 0.0,
            "Realized ($M)": 0.0 if status == "Active" else round(invested * 1.2, 6),
            "MOIC (x)": None,
            "ARR ($M)": arr if status == "Active" else None, "ARR Growth (YoY %)": 0.3 if status == "Active" else None,
            "Gross Margin (%)": 0.6 if status == "Active" else None, "Net Burn ($M/mo)": burn if status == "Active" else None,
            "Cash ($M)": round(burn * 20, 3) if status == "Active" else None, "Runway (mo)": None,
            "Headcount": 10 + i % 200 if status == "Active" else None,
        }
        row["MOIC (x)"] = round((row["Prior Mark ($M)"] + row["Realized ($M)"]) / invested, 6)
        row["Runway (mo)"] = 20.0 if status == "Active" else "-"
        positions.append(row)

    active = [p for p in positions if p["Status"] == "Active"]
    chosen = rng.sample(active, k)
    types = ["Priced Equity Round", "Convertible Note", "IPO", "Acquisition (Closed)", "Acquisition (Announced)",
             "Shutdown", "Secondary Sale", "Term Sheet Signed"]
    events: list[dict[str, Any]] = []
    for j, p in enumerate(chosen):
        t = types[j % len(types)]
        post, own = p["Latest Post-Money ($M)"], p["Ownership (FD %)"]
        d = date(2026, 7, 1) + timedelta(days=(j * 7) % 90)
        e: dict[str, Any] = {"date": d, "company": p["Company"], "event": t, "detail": "", "value": None,
                             "hc_investment": None, "ownership_after": None, "proceeds": None, "notes": "Synthetic event."}
        if t == "Priced Equity Round":
            e.update(detail="Series B", value=round(post * 1.4, 1), ownership_after=round(own * 0.95, 4))
        elif t == "Convertible Note":
            e.update(detail=f"$2.0M bridge note, ${round(post * 1.2, 1)}M valuation cap", hc_investment=0.3)
        elif t == "IPO":
            e.update(detail="Listed on Nasdaq", value=round(post * 2, 1), ownership_after=round(own * 0.9, 4))
        elif t == "Acquisition (Closed)":
            v = round(post * 1.5, 1)
            e.update(detail="All-cash acquisition", value=v, proceeds=round(own * v, 6))
        elif t == "Acquisition (Announced)":
            e.update(detail="Definitive agreement signed, all cash", value=round(post * 1.5, 1))
        elif t == "Shutdown":
            e.update(detail="Ceased operations", proceeds=0.0)
        elif t == "Secondary Sale":
            after = round(own / 2, 6)
            e.update(detail="HC sold half of its position", value=post, ownership_after=after,
                     proceeds=round((own - after) * post, 6))
        elif t == "Term Sheet Signed":
            e.update(detail="Series C term sheet", value=round(post * 1.3, 1))
        events.append(e)
    return positions, events


def build_synthetic(spec: dict[str, Any], out: Path) -> Path:
    positions, events = synthesise(spec)
    book = _load_base(BASE_WORKBOOK, BASE_ACTIVITY_SHEET)   # for the Field Definitions tab only
    act_header, act_rows = _activity_rows(events)
    names = spec.get("sheet_names") or {}
    wb = Workbook()
    _write_tab(wb, names.get("portfolio", "Portfolio"), book.portfolio_header, book.portfolio_header, positions, None, first=True)
    _write_tab(wb, names.get("activity", BASE_ACTIVITY_SHEET), act_header, act_header, act_rows, None)
    _write_others(wb, book.others)
    _save(wb, out)
    return out


def build_garbage(spec: dict[str, Any], out: Path) -> Path:
    kind = spec["garbage"]
    if kind == "csv":
        with out.open("w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(list(PORTFOLIO_COLUMNS))
            w.writerow(["Aravine", "Climate & Energy", "Fund I", "Series A", "Active", "2020-01-04", "2021-03-09",
                        50.8, 6, 0.136, 6.9, 0, 1.15, 3.4, 0.83, 0.49, 0.54, 18.1, 33.5, 28])
        return out
    if kind == "empty":
        out.write_bytes(b"")
        return out
    wb = Workbook()
    if kind == "no_portfolio":
        ws = wb.active
        ws.title = "Summary"
        ws.append(["This workbook has a cover sheet and an activity tab, but no Portfolio tab."])
        wa = wb.create_sheet(BASE_ACTIVITY_SHEET)
        wa.append(list(ACTIVITY_COLUMNS))
        wa.append([datetime(2026, 7, 6), "Dovelane Systems", "Priced Equity Round", "Series B", 103.8, None, 0.055, None, ""])
    elif kind == "headers_only":
        ws = wb.active
        ws.title = "Portfolio"
        ws.append(list(PORTFOLIO_COLUMNS))
        wa = wb.create_sheet(BASE_ACTIVITY_SHEET)
        wa.append(list(ACTIVITY_COLUMNS))
    else:
        raise ValueError(f"unknown garbage kind {kind!r}")
    _save(wb, out)
    return out


def build(spec: dict[str, Any]) -> Path | None:
    WORKBOOKS.mkdir(parents=True, exist_ok=True)
    out = WORKBOOKS / spec.get("workbook", f"{spec['name']}.xlsx")
    if out.stem != spec["name"] and (SCENARIOS / f"{out.stem}.yaml").exists():
        return None   # shares another scenario's workbook (e.g. 12b runs 12's file with different market data)
    if spec.get("garbage"):
        return build_garbage(spec, out)
    if spec.get("synthesize"):
        return build_synthetic(spec, out)
    return build_standard(spec, out)


def scenario_files(names: list[str] | None = None) -> list[Path]:
    files = sorted(SCENARIOS.glob("*.yaml"))
    if names:
        wanted = {n.removesuffix(".yaml") for n in names}
        files = [f for f in files if f.stem in wanted]
        missing = wanted - {f.stem for f in files}
        if missing:
            raise SystemExit(f"no such scenario(s): {sorted(missing)}")
    return files


def main(argv: list[str]) -> int:
    for f in scenario_files(argv):
        spec = load_scenario(f)
        out = build(spec)
        if out is None:
            print(f"{spec['name']:32s} -> (reuses workbooks/{spec['workbook']})")
        else:
            print(f"{spec['name']:32s} -> {out.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
