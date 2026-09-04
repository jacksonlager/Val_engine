"""The gauntlet: run every scenario in `scenarios/` through the real pipeline, compare the
result to the scenario's `expect:` block, and write report.md + report.json.

    python training/run_gauntlet.py                 # everything; exit 1 if any check fails
    python training/run_gauntlet.py 02_event_type_typos 20_garbage_csv

Each scenario runs in isolation: a temp copy of `rules/` (plus a generated `2026Q4.yaml`
when the policy is 2026Q4), the committed workbook, no override ledger, no proposals,
adjudication off, a pinned `generated_at`. The runner never crashes on a failing scenario:
an `IngestError` is the expected outcome for `ingest_ok: false`; any other exception is a
failure that carries the traceback's last line.

`tests/test_gauntlet.py` imports `run_scenario` from here so pytest and this script agree.
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
import time
import traceback
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import yaml

TRAINING = Path(__file__).resolve().parent
ROOT = TRAINING.parent
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))
if str(TRAINING) not in sys.path:
    sys.path.insert(0, str(TRAINING))

from generate import SCENARIOS, WORKBOOKS, load_scenario, scenario_files  # noqa: E402
from hc_valuation import pipeline  # noqa: E402
from hc_valuation.config import load_config  # noqa: E402
from hc_valuation.ingest.reader import IngestError, read_workbook  # noqa: E402

REPORT_MD = TRAINING / "report.md"
REPORT_JSON = TRAINING / "report.json"

# The Q4 policy a snapshot-based scenario runs under. Written into the scenario's temp rules
# dir at run time — never into the real rules/.
Q4_POLICY: dict[str, Any] = {
    "policy_version": "2026Q4-0.1",
    "inherits": "2026Q3",
    "quarter": {"label": "Q4 2026", "measurement_date": date(2026, 12, 31), "prior_close": date(2026, 9, 30),
                "window_start": date(2026, 10, 1), "window_end": date(2026, 12, 31)},
}


# ----------------------------------------------------------------------------- results

@dataclass
class Check:
    name: str
    expected: Any
    actual: Any
    ok: bool

    def line(self) -> str:
        return f"{'✓' if self.ok else '✗'} {self.name}: expected {_fmt(self.expected)}, got {_fmt(self.actual)}"


@dataclass
class ScenarioReport:
    name: str
    description: str
    checks: list[Check] = field(default_factory=list)
    seconds: float = 0.0
    crashed: str | None = None

    @property
    def failures(self) -> list[Check]:
        return [c for c in self.checks if not c.ok]

    @property
    def ok(self) -> bool:
        return not self.failures and self.crashed is None

    def add(self, name: str, expected: Any, actual: Any, ok: bool) -> None:
        self.checks.append(Check(name, expected, actual, bool(ok)))

    def summary(self) -> str:
        n_ok = sum(1 for c in self.checks if c.ok)
        status = "PASS" if self.ok else "FAIL"
        crash = f"  (crash: {self.crashed})" if self.crashed else ""
        return f"{status:4s} {self.name:34s} {n_ok:3d}/{len(self.checks):<3d} checks  {self.seconds:6.2f}s{crash}"


def _fmt(v: Any) -> str:
    if isinstance(v, float):
        return f"{v:.4f}"
    if isinstance(v, (list, tuple, set, frozenset)):
        return "[" + ", ".join(str(x) for x in (sorted(v) if isinstance(v, (set, frozenset)) else v)) + "]"
    return str(v)


# ----------------------------------------------------------------------------- running

def _prepare(spec: dict[str, Any], tmp: Path) -> tuple[pipeline.RunPaths, str]:
    rules = tmp / "rules"
    shutil.copytree(ROOT / "rules", rules)
    policy = str(spec.get("policy", "2026Q3"))
    if policy == "2026Q4" and not (rules / "2026Q4.yaml").exists():
        (rules / "2026Q4.yaml").write_text(yaml.safe_dump(Q4_POLICY, sort_keys=False))
    policy_path = rules / f"{policy}.yaml"
    if not policy_path.exists():
        raise FileNotFoundError(f"policy {policy!r} not in rules/ and not generated")
    overrides = spec.get("policy_overrides")
    if overrides:
        # A scenario-local policy (custom_rules, a tolerance) written into the temp rules dir as a
        # child of the named policy — the real rules/ is never touched.
        child = {"inherits": policy, **overrides}
        child.setdefault("policy_version", f"{policy}-scenario")
        policy_path = rules / f"{policy}_{spec['name']}.yaml"
        policy_path.write_text(yaml.safe_dump(child, sort_keys=False))

    src_wb = WORKBOOKS / spec.get("workbook", f"{spec['name']}.xlsx")
    if not src_wb.exists():
        raise FileNotFoundError(f"workbook missing: {src_wb} (run `python training/generate.py {spec['name']}`)")
    workbook = tmp / src_wb.name
    shutil.copy(src_wb, workbook)
    sidecar = src_wb.with_name(src_wb.stem + ".open_items_carry.yaml")
    carry = tmp / "open_items_carry.yaml"
    if sidecar.exists():
        shutil.copy(sidecar, carry)
    # data/mock_responses lives in the repo root; the stub connectors fall back to it.
    paths = pipeline.RunPaths(root=tmp, policy=policy_path, workbook=workbook, overrides=tmp / "overrides.yaml",
                              proposals_dir=tmp / "proposals", precedent=tmp / "precedent.yaml", open_items_carry=carry)
    return paths, policy


class _MarketPatch:
    """`market: {drop_quotes: [...]}` — remove quotes after the connectors assemble them, so a
    scenario can ask what the engine does when a listed position has no measurement-date price."""

    def __init__(self, drop: list[str]) -> None:
        self.drop = set(drop)
        self._orig = None

    def __enter__(self) -> "_MarketPatch":
        import hc_valuation.connectors as connectors
        self._orig = connectors.assemble_market_data
        orig = self._orig
        drop = self.drop

        def patched(cfg, root, snapshot, feed, provider=None, **kw):
            assembled = orig(cfg, root, snapshot, feed, provider=provider, **kw)
            md, label = assembled
            quotes = {k: v for k, v in md.quotes.items() if k not in drop}
            return connectors.MarketAssembly(market=md.model_copy(update={"quotes": quotes}), label=label,
                                             report=dict(getattr(assembled, "report", None) or {}))

        connectors.assemble_market_data = patched
        return self

    def __exit__(self, *exc: Any) -> None:
        import hc_valuation.connectors as connectors
        connectors.assemble_market_data = self._orig


def run_scenario(path: Path) -> ScenarioReport:
    spec = load_scenario(path)
    rep = ScenarioReport(name=spec["name"], description=str(spec.get("description", "")).strip())
    expect = spec.get("expect") or {}
    t0 = time.perf_counter()
    with tempfile.TemporaryDirectory(prefix=f"gauntlet_{spec['name']}_") as td:
        tmp = Path(td)
        try:
            paths, _policy = _prepare(spec, tmp)
        except Exception as exc:  # noqa: BLE001 - setup failure is a scenario failure, never a crash of the runner
            rep.crashed = f"setup: {exc}"
            rep.seconds = time.perf_counter() - t0
            return rep

        result = None
        error: BaseException | None = None
        drop = (spec.get("market") or {}).get("drop_quotes") or []
        try:
            with _MarketPatch(drop):
                cfg = load_config(paths.policy)
                result = pipeline.execute(paths, generated_at=datetime.combine(cfg.quarter.measurement_date, datetime.min.time()),
                                          adjudicate=False, provider="stub")
        except IngestError as exc:
            error = exc
        except Exception as exc:  # noqa: BLE001
            last = traceback.format_exception_only(type(exc), exc)[-1].strip()
            rep.crashed = last
            rep.add("no_crash", "no exception", last, False)
        rep.seconds = time.perf_counter() - t0
        if rep.crashed is None:
            try:
                _check(rep, spec, expect, paths, result, error, tmp)
            except Exception as exc:  # noqa: BLE001 - a bug in a check must not take the runner down
                last = traceback.format_exception_only(type(exc), exc)[-1].strip()
                rep.crashed = f"checker: {last}"
                rep.add("checker", "checks evaluated", last, False)
    return rep


# ----------------------------------------------------------------------------- checks

def _contains_any(text: str, needles: Any) -> bool:
    if needles is None:
        return True
    if isinstance(needles, str):
        needles = [needles]
    low = text.lower()
    return any(str(n).lower() in low for n in needles)


def _check(rep: ScenarioReport, spec: dict[str, Any], expect: dict[str, Any], paths: pipeline.RunPaths,
           result: Any, error: BaseException | None, tmp: Path) -> None:
    ingest_ok = expect.get("ingest_ok", True)
    if ingest_ok is False:
        rep.add("ingest_ok", False, error is None, error is not None)
        if error is not None:
            msg = str(error)
            rep.add("error_contains", expect.get("error_contains"), msg[:160], _contains_any(msg, expect.get("error_contains")))
            rep.add("error_is_not_a_traceback", "message names the problem", msg[:80],
                    "Traceback" not in msg and "line " not in msg.lower()[:40])
        return
    if ingest_ok == "any":
        if error is not None:
            msg = str(error)
            rep.add("ingest_error_names_the_cell", expect.get("error_contains"), msg[:160],
                    _contains_any(msg, expect.get("error_contains")))
            return
    else:
        rep.add("ingest_ok", True, error is None, error is None)
        if error is not None:
            rep.add("ingest_error", "none", str(error)[:200], False)
            return
    if result is None:
        return

    run = result.run
    if expect.get("max_seconds") is not None:
        rep.add("max_seconds", f"< {expect['max_seconds']}s", round(rep.seconds, 2), rep.seconds < float(expect["max_seconds"]))
    if expect.get("no_crash", True):
        rep.add("no_crash", "no exception", "no exception", True)

    # ---- validation issues
    issues = list(run.validation)
    ids = {v.rule_id for v in issues}
    vids = expect.get("validation_ids") or {}
    for i in vids.get("must", []) or []:
        rep.add(f"validation must include {i}", i, sorted(ids), i in ids)
    for i in vids.get("must_not", []) or []:
        rep.add(f"validation must not include {i}", f"no {i}", sorted(ids), i not in ids)
    for m in expect.get("validation_must", []) or []:
        hits = _matching_issues(issues, m)
        want = m.get("count")
        label = f"validation {m.get('id', '*')}" + (f" on {m['company']}" if m.get("company") else "") \
            + (f" containing {m['contains']!r}" if m.get("contains") else "")
        if want is not None:
            rep.add(label + " count", want, len(hits), len(hits) == int(want))
        else:
            rep.add(label, "≥ 1", len(hits), len(hits) >= 1)
    for m in expect.get("validation_must_not", []) or []:
        hits = _matching_issues(issues, m)
        label = f"no validation {m.get('id', '*')}" + (f" on {m['company']}" if m.get("company") else "")
        rep.add(label, 0, len(hits), len(hits) == 0)
    if "validation_blocking" in expect:
        blocking = sorted({f"{v.rule_id}:{v.company or v.sheet}" for v in issues if v.blocking})
        rep.add("validation_blocking", bool(expect["validation_blocking"]), blocking or False,
                bool(blocking) == bool(expect["validation_blocking"]))
    if expect.get("anywhere_must"):
        flag_ids = {f.rule_id for c in run.companies for f in c.flags}
        for i in expect["anywhere_must"]:
            rep.add(f"{i} raised (validation or flag)", i, sorted((ids | flag_ids) & {i}) or "absent", i in ids or i in flag_ids)

    # ---- totals / manifest
    totals = expect.get("totals") or {}
    if "events" in totals:
        try:
            _snap, feed = read_workbook(paths.workbook, result.config)
            n_events = len(feed.events)
        except Exception as exc:  # noqa: BLE001
            n_events = f"re-read failed: {exc}"
        rep.add("totals.events", totals["events"], n_events, n_events == totals["events"])
    if "positions" in totals:
        rep.add("totals.positions", totals["positions"], run.totals.positions, run.totals.positions == totals["positions"])
    if "dispositions" in totals:
        rep.add("totals.dispositions", totals["dispositions"], run.totals.dispositions,
                run.totals.dispositions == totals["dispositions"])
    man = expect.get("manifest") or {}
    for k, v in man.items():
        actual = getattr(run.manifest, k, None)
        actual = actual.isoformat() if isinstance(actual, (date, datetime)) else actual
        rep.add(f"manifest.{k}", v, actual, str(actual) == str(v))

    # ---- companies
    by = run.by_company()
    for name, exp in (expect.get("companies") or {}).items():
        exp = exp or {}
        if exp.get("present") is False:
            rep.add(f"{name}: absent from the run", "absent", "present" if name in by else "absent", name not in by)
            continue
        c = by.get(name)
        if c is None:
            rep.add(f"{name}: present in the run", "present", "absent", False)
            continue
        _check_company(rep, name, exp, c)

    # ---- snapshot round-trip (year-end rollover etc.)
    snap = expect.get("snapshot")
    if snap:
        _check_snapshot(rep, snap, result, paths, tmp)


def _matching_issues(issues: list, m: dict[str, Any]) -> list:
    out = []
    for v in issues:
        if m.get("id") and v.rule_id != m["id"]:
            continue
        if m.get("company") and (v.company or "") != m["company"]:
            continue
        if m.get("contains") and not _contains_any(v.message, m["contains"]):
            continue
        if "blocking" in m and bool(v.blocking) != bool(m["blocking"]):
            continue
        out.append(v)
    return out


def _check_company(rep: ScenarioReport, name: str, exp: dict[str, Any], c: Any) -> None:
    chain = [s.rule_id for s in c.steps]
    flags = {f.rule_id for f in c.flags}
    tol = float(exp.get("tol", 0.01))
    if "disposition" in exp:
        rep.add(f"{name}: disposition", exp["disposition"], c.disposition.value, c.disposition.value == exp["disposition"])
    if "disposition_in" in exp:
        rep.add(f"{name}: disposition in", list(exp["disposition_in"]), c.disposition.value, c.disposition.value in exp["disposition_in"])
    if "rules" in exp:
        rep.add(f"{name}: rule chain", list(exp["rules"]), chain, chain == list(exp["rules"]))
    for r in exp.get("rules_contains", []) or []:
        rep.add(f"{name}: chain contains {r}", r, chain, r in chain)
    for r in exp.get("rules_must_not", []) or []:
        rep.add(f"{name}: chain must not contain {r}", f"no {r}", chain, r not in chain)
    if "proposed" in exp:
        want = float(exp["proposed"])
        rep.add(f"{name}: proposed mark", want, c.proposed_mark, abs(c.proposed_mark - want) <= tol)
    for f in exp.get("flags_must", []) or []:
        rep.add(f"{name}: flag {f} present", f, sorted(flags), f in flags)
    for f in exp.get("flags_must_not", []) or []:
        rep.add(f"{name}: flag {f} absent", f"no {f}", sorted(flags), f not in flags)
    for attr, key in (("status_after", "status"), ("stage", "stage"), ("fund", "fund"), ("sector", "sector")):
        if key in exp:
            actual = getattr(c, attr)
            actual = actual.value if hasattr(actual, "value") else actual
            rep.add(f"{name}: {key}", exp[key], actual, actual == exp[key])
    for key in ("listed", "fv_level"):
        if key in exp:
            rep.add(f"{name}: {key}", exp[key], getattr(c, key), getattr(c, key) == exp[key])
    for key in ("realized_quarter", "ownership_after", "invested_after", "note_at_cost", "equity_mark", "latest_post_money"):
        if key in exp:
            actual = float(getattr(c, key))
            rep.add(f"{name}: {key}", float(exp[key]), actual, abs(actual - float(exp[key])) <= tol)
    items = list(c.open_items)
    for want in exp.get("open_items_must", []) or []:
        want = {"kind": want} if isinstance(want, str) else want
        hit = [o for o in items if o.kind.value == want["kind"]
               and all(getattr(o, k) == v for k, v in want.items() if k != "kind")]
        rep.add(f"{name}: open item {want}", "present", [(o.kind.value, o.age_quarters, o.escalated) for o in items], bool(hit))
    for kind in exp.get("open_items_must_not", []) or []:
        rep.add(f"{name}: no open item {kind}", f"no {kind}", [o.kind.value for o in items],
                not any(o.kind.value == kind for o in items))
    if "open_items_count" in exp:
        rep.add(f"{name}: open items count", exp["open_items_count"], len(items), len(items) == int(exp["open_items_count"]))
    for k, v in (exp.get("alternative_marks") or {}).items():
        actual = c.alternative_marks.get(k)
        rep.add(f"{name}: alternative mark {k}", float(v), actual, actual is not None and abs(actual - float(v)) <= tol)


def _check_snapshot(rep: ScenarioReport, snap: dict[str, Any], result: Any, paths: pipeline.RunPaths, tmp: Path) -> None:
    from hc_valuation.export.snapshot import next_quarter_label, write_next_quarter_workbook
    try:
        label = next_quarter_label(result.config.quarter.label)
        out = write_next_quarter_workbook(result.run, paths.workbook, tmp / "next" / "next_quarter.xlsx", result.config)
        import openpyxl
        wb = openpyxl.load_workbook(out, read_only=True)
        sheets = list(wb.sheetnames)
        n_rows = sum(1 for r in wb["Portfolio"].iter_rows(min_row=2, values_only=True) if r and r[0])
        wb.close()
    except Exception as exc:  # noqa: BLE001
        rep.add("snapshot: emitted", "next-quarter workbook written", f"{type(exc).__name__}: {exc}"[:160], False)
        return
    if "next_label" in snap:
        rep.add("snapshot: next quarter label", snap["next_label"], label, label == snap["next_label"])
    if "activity_sheet" in snap:
        rep.add("snapshot: activity sheet", snap["activity_sheet"], sheets, snap["activity_sheet"] in sheets)
    if "portfolio_rows" in snap:
        rep.add("snapshot: portfolio rows", snap["portfolio_rows"], n_rows, n_rows == int(snap["portfolio_rows"]))
    rep.add("snapshot: sidecar written", "open_items_carry.yaml", "present" if (out.parent / "open_items_carry.yaml").exists() else "absent",
            (out.parent / "open_items_carry.yaml").exists())


# ----------------------------------------------------------------------------- reporting

def write_reports(reports: list[ScenarioReport], when: datetime) -> None:
    n_pass = sum(1 for r in reports if r.ok)
    n_checks = sum(len(r.checks) for r in reports)
    n_fail = sum(len(r.failures) for r in reports)
    lines = ["# Gauntlet report", "",
             f"Run at {when.isoformat(timespec='seconds')} — **{n_pass}/{len(reports)} scenarios green**, "
             f"{n_checks - n_fail}/{n_checks} checks passed.", "",
             "| Scenario | Status | Checks | Time |", "|---|---|---:|---:|"]
    for r in reports:
        ok = sum(1 for c in r.checks if c.ok)
        status = "✓ pass" if r.ok else ("✗ crash" if r.crashed else "✗ fail")
        lines.append(f"| [{r.name}](#{r.name.replace('_', '-')}) | {status} | {ok}/{len(r.checks)} | {r.seconds:.2f}s |")
    lines.append("")
    lines.append("## Failures")
    lines.append("")
    any_fail = False
    for r in reports:
        if r.crashed:
            any_fail = True
            lines.append(f"- **{r.name}** crashed: `{r.crashed}`")
        for c in r.failures:
            any_fail = True
            lines.append(f"- **{r.name}** — {c.name}: expected `{_fmt(c.expected)}`, got `{_fmt(c.actual)}`")
    if not any_fail:
        lines.append("None.")
    lines.append("")
    for r in reports:
        lines.append(f"## {r.name}")
        lines.append("")
        if r.description:
            lines.append(r.description)
            lines.append("")
        if r.crashed:
            lines.append(f"**Crashed:** `{r.crashed}`")
            lines.append("")
        lines.append("| Check | Expected | Actual | |")
        lines.append("|---|---|---|:-:|")
        for c in r.checks:
            lines.append(f"| {_md(c.name)} | {_md(_fmt(c.expected))} | {_md(_fmt(c.actual))} | {'✓' if c.ok else '✗'} |")
        lines.append("")
    REPORT_MD.write_text("\n".join(lines))
    payload = {
        "generated_at": when.isoformat(timespec="seconds"),
        "scenarios": len(reports), "passed": n_pass, "checks": n_checks, "failed_checks": n_fail,
        "results": [{
            "name": r.name, "ok": r.ok, "seconds": round(r.seconds, 3), "crashed": r.crashed,
            "checks": [{"name": c.name, "expected": _json(c.expected), "actual": _json(c.actual), "ok": c.ok} for c in r.checks],
        } for r in reports],
    }
    REPORT_JSON.write_text(json.dumps(payload, indent=2, ensure_ascii=False))


def _md(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ")[:300]


def _json(v: Any) -> Any:
    if isinstance(v, (set, frozenset)):
        return sorted(v)
    if isinstance(v, (list, tuple)):
        return [_json(x) for x in v]
    if isinstance(v, dict):
        return {str(k): _json(x) for k, x in v.items()}
    if isinstance(v, (str, int, float, bool)) or v is None:
        return v
    return str(v)


def main(argv: list[str]) -> int:
    files = scenario_files(argv)
    reports: list[ScenarioReport] = []
    for f in files:
        rep = run_scenario(f)
        reports.append(rep)
        print(rep.summary(), flush=True)
    write_reports(reports, datetime.now().replace(microsecond=0))
    n_pass = sum(1 for r in reports if r.ok)
    n_checks = sum(len(r.checks) for r in reports)
    n_fail = sum(len(r.failures) for r in reports)
    print(f"\n{n_pass}/{len(reports)} scenarios green; {n_checks - n_fail}/{n_checks} checks passed. "
          f"Report: {REPORT_MD.relative_to(ROOT)}")
    return 0 if n_pass == len(reports) else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
